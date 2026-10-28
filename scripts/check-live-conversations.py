"""Opt-in live API check. Uses existing model credentials and isolated local data."""
import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
parser = argparse.ArgumentParser()
parser.add_argument("--config", type=Path, required=True)
parser.add_argument("--roles", default="1,2,3,4,5,6,talk")
parser.add_argument("--output", type=Path)
args = parser.parse_args()
out = (args.output or ROOT / "backend/workspace/live-conversation-check" / datetime.now().strftime("%Y%m%d-%H%M%S")).resolve()
out.mkdir(parents=True, exist_ok=True)
os.environ["OFFICE_DB_PATH"] = str(out / "office.db")

import agent_config
agent_config.CONFIG_PATH = str(args.config.resolve())
import forge
forge.WORKSPACE_DIR = str(out / "workspace")
forge.DELIVERABLES_DIR = str(out / "deliverables")
forge.CODE_ROOT = str(out / "workspace/code")
forge.CODE_DIR = forge.CODE_ROOT
import llm_usage
llm_usage.USAGE_FILE = str(out / "usage.json")
import main
import storage
import llm
from docx import Document
from fastapi.testclient import TestClient

report = {"output": str(out), "checks": [], "started": datetime.now().isoformat()}
history = {}


def log(message):
    print(message, flush=True)


def save():
    report["usage"] = {"calls": llm.usage_total_calls(), "cost": llm.usage_total_cost()}
    (out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def post(client, path, body):
    response = client.post(path, json=body)
    if response.status_code != 200:
        raise AssertionError(f"{path}: HTTP {response.status_code}: {response.text[:400]}")
    return response.json()


def turn(client, scope, message, sprint_id=""):
    roster = 1 if scope == "talk" else int(scope)
    messages = history.setdefault(scope, [])
    log(f"[{scope}] sending: {message}")
    result = post(client, "/api/chat/turn", {
        "roster": roster, "mode": "talk" if scope == "talk" else "solo",
        "session_id": f"live-check-{scope}", "message": message,
        "history": messages, "sprint_id": sprint_id,
    })
    messages.append({"role": "user", "content": message})
    messages.append({"role": "assistant", "content": result.get("reply", "已提交任务")})
    return result


def artifact_text(client, artifact):
    assert not artifact.get("degraded"), f"Degraded output: {artifact.get('degraded_reason')}"
    assert artifact.get("source") == "model", f"Not real model output: {artifact.get('source')}"
    assert artifact.get("file_name"), f"No actual file: {artifact.get('degraded_reason')}"
    aid = artifact["id"]
    response = client.get(f"/api/artifacts/{aid}/download")
    assert response.status_code == 200 and response.content.startswith(b"PK"), "DOCX download failed"
    path = forge.safe_deliverable_path(artifact["file_name"])
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert client.get(f"/api/artifacts/{aid}/preview").status_code == 200, "preview failed"
    return text


def contains(text, *parts):
    compact = re.sub(r"\s+", "", text)
    for part in parts:
        assert part in compact, f"Missing {part!r} in {text!r}"


with TestClient(main.app) as client:
    for scope in args.roles.split(","):
        check = {"scope": scope, "model": llm.effective_model(1 if scope == "talk" else int(scope)), "status": "running"}
        report["checks"].append(check)
        start = time.monotonic()
        try:
            prompt = "请生成一份简短中文Word通知，标题为联调通知。通知说明：2026年10月29日14:00在A301会议室举行联调，预算800元，联系人王明。保持简洁，不要封面。"
            first_result = turn(client, scope, prompt)
            sid = ""
            if scope == "talk":
                assert first_result.get("intent") == "requirement" and not first_result.get("handled")
                started = post(client, "/api/sprint/start", {
                    "requirement": first_result["requirement"], "execution_style": "quick", "task_kind": "file", "budget": 2,
                    "save_dir": str(out / "deliverables"),
                })
                sid = started["sprint_id"]
                deadline = time.monotonic() + 600
                while main.runtime_state.running and time.monotonic() < deadline:
                    time.sleep(1)
                assert not main.runtime_state.running, "Team generation timed out"
                products = storage.list_artifacts(sid)
            else:
                assert first_result.get("handled"), first_result
                products = first_result.get("artifacts", [])
            assert len(products) == 1, f"Expected one product, got {len(products)}"
            first = products[0]
            original = artifact_text(client, first)
            contains(original, "14:00", "A301", "800", "王明")
            log(f"[{scope}] v1 actual DOCX verified")
            second_result = turn(client, scope, "把时间改成2026年10月30日15:30，其余内容包括地点、预算、联系人保持不变。", sid)
            second = second_result.get("artifacts", [None])[0]
            assert second, second_result
            contains(artifact_text(client, second), "15:30", "A301", "800", "王明")
            assert second["parent_id"] == first["id"] and second["version"] == 2
            log(f"[{scope}] v2 time change and retained content verified")
            third_result = turn(client, scope, "再把预算改成1200元，其他保持不变。", sid)
            third = third_result.get("artifacts", [None])[0]
            assert third, third_result
            contains(artifact_text(client, third), "15:30", "A301", "1200", "王明")
            assert third["parent_id"] == second["id"] and third["version"] == 3
            log(f"[{scope}] v3 retained time change and updated budget verified")
            restored_result = turn(client, scope, "恢复到v1的内容。", sid)
            restored = restored_result.get("artifacts", [None])[0]
            assert restored, restored_result
            assert artifact_text(client, restored) == original, "Rollback changed original content"
            assert artifact_text(client, first) == original, "Original file was mutated"
            assert restored["version"] == 4
            check.update(status="passed", artifact_ids=[a["id"] for a in (first, second, third, restored)])
            log(f"[{scope}] PASS four-turn live model flow, downloads, previews and preserved original")
        except Exception as exc:
            check.update(status="failed", error=f"{type(exc).__name__}: {str(exc)[:1200]}")
            log(f"[{scope}] FAIL: {check['error']}")
        finally:
            check["seconds"] = round(time.monotonic() - start, 2)
            save()
report["finished"] = datetime.now().isoformat()
save()
log(f"REPORT: {out / 'report.json'}")
sys.exit(1 if any(c["status"] != "passed" for c in report["checks"]) else 0)
