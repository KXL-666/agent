"""Completion regressions, using real temporary files and mocked model/runtime calls."""
import copy
import json
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

import artifacts
import artifact_quality
import forge
import llm
import main
import project_agent
import revision_chunks as chunks
import revision_intent as intent
import revision_jobs
import revision_diff
import revision_outputs as outputs
import revision_service as service
import revision_visual as visual
import storage
import verification
import test_revision_sessions as fixtures
import test_revision_extensions as extensions


class RevisionCompletionTest(unittest.IsolatedAsyncioTestCase):
    setUp = fixtures.RevisionSessionTest.setUp
    asyncTearDown = fixtures.RevisionSessionTest.asyncTearDown
    accepted = extensions.RevisionExtensionsTest.accepted
    generate_file = extensions.RevisionExtensionsTest.generate_file

    def test_explicit_natural_language_routing(self):
        for message, expected in [
            ('再配套一份PPT和Excel', ('add', ['pptx', 'xlsx'])),
            ('把Word转换成PDF', ('convert', ['pdf'])),
            ('先不要生成，只解释怎么转成PDF', ('discuss', [])),
            ('再把PPT标题改短', ('revise', [])),
            ('把联调时间改为16:00', ('revise', [])),
        ]:
            with self.subTest(message=message):
                self.assertEqual(intent.infer(message), expected)
        with self.assertRaises(HTTPException):
            intent.infer('转成PPT和PDF')

    def test_contract_repair_requires_all_exact_quotes_and_preserves_retirement_scope(self):
        feedback = '这是联调测试，请通知同事准备账号'
        plan = {'target_kind': 'docx', 'audience': '同事', 'purpose': '通知',
                'superseded_ids': [], 'requirements': [], 'missing_info': [], 'allowed_assumptions': []}
        repaired = {**plan, 'requirements': [
            {'source': 'user', 'source_quote': quote, 'acceptance_method': '读取真实文件核对该条原文含义'}
            for quote in ('这是联调测试', '请通知同事准备账号')]}
        with patch.object(main, 'chat_json', side_effect=[plan, repaired]) as model:
            effective = service.evolve_contract(self.original, feedback)
        self.assertEqual(model.call_count, 2)
        self.assertEqual(effective['status'], 'ready')
        self.assertEqual(json.loads(model.call_args_list[0].args[1])['exact_feedback_fragments'],
                         ['这是联调测试', '请通知同事准备账号'])
        with patch.object(main, 'chat_json', side_effect=[plan, {**repaired, 'superseded_ids': ['R1']}]), \
             self.assertRaises(HTTPException):
            service.evolve_contract(self.original, feedback)

    def test_normalized_model_source_objects_require_exact_original_quotes(self):
        plan = {'target_kind': 'docx', 'audience': '同事', 'purpose': '通知', 'superseded_ids': [],
            'requirements': [{'source': {'type': 'user', 'quote': '仅使用当前通知中的事实'},
                              'acceptance_method': '逐项核对原通知'}],
            'missing_info': [], 'allowed_assumptions': []}
        with patch.object(main, 'chat_json', return_value=plan) as model:
            effective = service.evolve_contract(self.original, '仅使用当前通知中的事实')
        self.assertEqual(model.call_count, 1)
        self.assertEqual(effective['status'], 'ready')
        bad = copy.deepcopy(plan); bad['requirements'][0]['source']['quote'] = '模型改写的句子'
        with patch.object(main, 'chat_json', return_value=bad), self.assertRaises(HTTPException):
            service.evolve_contract(self.original, '仅使用当前通知中的事实')

    def test_contract_repair_does_not_accept_incomplete_model_annotations(self):
        plan = {'target_kind': 'docx', 'audience': '同事', 'purpose': '通知',
                'superseded_ids': [], 'requirements': [], 'missing_info': [], 'allowed_assumptions': []}
        with patch.object(main, 'chat_json', return_value=plan) as model, self.assertRaises(HTTPException):
            service.evolve_contract(self.original, '这是联调测试，请通知同事准备账号')
        self.assertEqual(model.call_count, 2)

    def test_compound_negation_does_not_require_a_toc_or_override_other_checks(self):
        clauses = [{'text': '不需要封面、目录和插图'}]
        items = [{'status': 'passed'}]
        verification._apply_mechanical_toc_checks(items, clauses, {'toc_present': False})
        self.assertEqual(items[0]['status'], 'passed')
        items = [{'status': 'failed', 'problem': '有独立封面'}]
        verification._apply_mechanical_toc_checks(items, clauses, {'toc_present': False})
        self.assertEqual(items[0]['status'], 'failed')
        for text in ('无目录', '不要目录', '不需要封面、目录和插图'):
            self.assertFalse(artifact_quality._affirmed(text, r'目录'))
        self.assertTrue(artifact_quality._affirmed('不需要封面，但需要目录', r'目录'))

    def test_retired_clause_keeps_fact_provenance_without_remaining_active(self):
        quote = '2026年10月29日14:00在A301会议室联调'
        original = copy.deepcopy(self.original)
        original['regen'] = {'revision_contract': {'requirements': [
            {'requirement_id': 'R1', 'text': quote, 'source': {'type': 'user', 'quote': quote},
             'status': 'ready', 'acceptance_method': '核对通知时间地点'}]}}
        plan = {'target_kind': 'docx', 'audience': '同事', 'purpose': '通知', 'superseded_ids': ['R1'],
            'requirements': [{'source': 'user', 'source_quote': '改为16:00', 'acceptance_method': '核对新时间'}],
            'missing_info': [], 'allowed_assumptions': []}
        with patch.object(main, 'chat_json', return_value=plan):
            current = service.evolve_contract(original, '改为16:00')
        self.assertEqual(current['raw_requirement'], '改为16:00')
        self.assertEqual(current['prior_user_quotes'], [quote])
        material = verification.source_material(current)
        self.assertIn('2026年10月29日', material)
        self.assertIn('旧要求已失效', material)
        self.assertEqual(artifact_quality.requirement_scope(material), '改为16:00')

    def test_text_replacement_is_not_counted_as_a_font_change(self):
        path = Path(forge.safe_deliverable_path(self.original['file_name']))
        document = Document(path)
        document.paragraphs[0].runs[0].text = '联调时间16:00'
        name = 'changed-text.docx'; document.save(forge.output_path(name))
        item = artifacts.add_artifact('final', '时间修订', 3, '', file_name=name, sprint_id=self.sid)
        result = revision_diff.compare(item['id'], self.original['id'])
        self.assertEqual(result['content_changes'], 1)
        self.assertEqual(result['format_changes'], 0)
        document.paragraphs[0].runs[0].bold = True
        document.save(forge.safe_deliverable_path(item['file_name']))
        self.assertEqual(revision_diff.compare(item['id'], self.original['id'])['format_changes'], 1)

    async def test_auto_request_resolves_once_and_rejects_changed_retry(self):
        request = {'request_id': 'auto-add', 'expected_artifact_id': self.original['id'],
                   'feedback': '再配套一份PPT和Excel'}
        url = f'/api/revision/session/{self.session["id"]}/message'
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://test') as client:
            with patch.object(revision_jobs, 'dispatch'):
                first = await client.post(url, json=request)
            self.assertEqual(first.status_code, 202, first.text)
            operation = first.json()['operation']
            self.assertEqual(operation['action'], 'add')
            self.assertEqual(operation['payload']['target_kinds'], ['pptx', 'xlsx'])
            storage.finish_revision_operation(operation['id'], artifact=self.original, summary='done')
            with patch.object(intent, 'infer', side_effect=AssertionError('retry must not resolve again')), \
                 patch.object(revision_jobs, 'dispatch') as dispatch:
                retry = await client.post(url, json=request)
            self.assertEqual(retry.status_code, 202)
            self.assertEqual(retry.json()['operation']['id'], operation['id'])
            dispatch.assert_not_called()
            conflict = await client.post(url, json={**request, 'feedback': '转成PDF'})
            self.assertEqual(conflict.status_code, 409)

    async def test_duplicate_add_is_rejected_before_model_job_is_created(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://test') as client:
            with patch.object(revision_jobs, 'dispatch') as dispatch:
                response = await client.post(f'/api/revision/session/{self.session["id"]}/message', json={
                    'request_id': 'duplicate-word', 'expected_artifact_id': self.original['id'],
                    'feedback': '再配套一份Word'})
            self.assertEqual(response.status_code, 409)
            dispatch.assert_not_called()
            self.assertFalse(storage.revision_operations(self.session['id']))

    def test_incomplete_bundle_files_do_not_appear_as_completed_history_versions(self):
        pending = artifacts.add_artifact('final', '尚未完成的配套', 3, '', file_name=self.original['file_name'],
            parent_id=self.original['id'], sprint_id=self.sid, version=2, regen={'bundle_pending': True})
        snapshot = revision_jobs.session_snapshot(self.session['id'])
        self.assertNotIn(pending['id'], [version['id'] for version in snapshot['versions']])

    def test_companion_contract_scopes_source_format_but_keeps_facts(self):
        clauses = [
            {'text': '请生成一份简短中文Word通知'},
            {'text': '预算1500元'},
            {'text': '介绍Word的用途'},
            {'text': '标题《联调通知》', 'acceptance_method': '检查Word Title样式'},
        ]
        for i, item in enumerate(clauses, 1):
            item.update(requirement_id=f'R{i}', source={'type': 'user', 'quote': item['text']}, status='ready')
        current = service.format_contract({'requirements': clauses}, 'pptx')
        self.assertEqual([c['text'] for c in current['requirements']], ['预算1500元', '介绍Word的用途', '标题《联调通知》'])
        self.assertEqual(current['format_scoped_sources'][0]['quote'], clauses[0]['text'])
        self.assertIn('PPT实际可见', current['requirements'][-1]['acceptance_method'])

    async def test_repairing_companion_recovers_bundle_after_former_anchor_failed(self):
        child = artifacts.add_artifact('final', '未通过配套', 3, '', file_name=self.original['file_name'],
            sprint_id=self.sid, parent_id=self.original['id'], quality_status='failed')
        records = [{'slot': 'source', 'artifact_id': self.original['id']},
                   {'slot': 'companion', 'artifact_id': child['id'], 'sync_from': 'source'}]
        first = await outputs._seal(self.original, records, [], '配套验收失败')
        self.assertEqual(first['artifact']['quality_status'], 'failed')
        fixed = artifacts.add_artifact('final', '已修正配套', 3, '', file_name=self.original['file_name'],
            sprint_id=self.sid, parent_id=child['id'], quality_status='passed')
        records[1]['artifact_id'] = fixed['id']
        second = await outputs._seal(fixed, records, [], '仅修复配套')
        self.assertEqual(second['artifact']['quality_status'], 'passed')
        self.assertEqual(second['outputs'][0]['artifact']['quality_status'], 'passed')
        self.assertEqual(storage.get_artifact(self.original['id'])['quality_status'], 'failed')

    async def semantic_pair(self):
        with patch.object(forge, 'forge_file', side_effect=self.generate_file), \
             patch.object(service, '_validate', side_effect=self.accepted):
            return (await outputs.edit(self.original, '新增配套PPT', self.contract,
                action='add', target_kind='pptx', operation_id='pair'))['artifact']

    async def test_semantic_sync_uses_new_source_and_preserves_companion_contract(self):
        anchor = await self.semantic_pair()
        old = next(o['artifact'] for o in outputs.output_snapshot(anchor) if o['kind'] == 'pptx')
        contract = {**self.contract, 'raw_requirement': 'PPT保留三页结构', 'target_kind': 'pptx',
            'requirements': [{'requirement_id': 'R1', 'text': 'PPT保留三页结构', 'status': 'ready',
                              'source': {'type': 'user', 'quote': 'PPT保留三页结构'},
                              'acceptance_method': '检查三页结构'}]}
        captured = []

        async def rewrite(item, effective, instruction):
            captured.append((item, effective, instruction))
            name = self.generate_file('pptx', '', '', 3)
            return {'file_name': name}

        with patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '16:00'}), \
             patch.object(service, 'evolve_contract', return_value=copy.deepcopy(contract)), \
             patch.object(service, '_rewrite', side_effect=rewrite), \
             patch.object(service, '_validate', side_effect=self.accepted):
            result = await outputs.edit(anchor, '时间改为16:00', {**self.contract, 'latest_feedback': '时间改为16:00'},
                artifact_id=self.original['id'], operation_id='sync', mode='targeted')
        self.assertEqual(len(captured), 1)
        self.assertEqual(captured[0][0]['id'], old['id'])
        self.assertIn('16:00', captured[0][2])
        self.assertEqual(captured[0][1]['raw_requirement'], 'PPT保留三页结构')
        snapshot = outputs.output_snapshot(result['artifact'])
        word = next(o['artifact'] for o in snapshot if o['kind'] == 'docx')
        ppt = next(o['artifact'] for o in snapshot if o['kind'] == 'pptx')
        self.assertEqual(ppt['regen']['sync_source_artifact_id'], word['id'])
        self.assertEqual(ppt['regen']['sync_source_sha256'], word['regen']['sha256'])
        self.assertIsNone(outputs.members(result['artifact'])[0]['sync_from'])

    async def test_semantic_sync_failure_preserves_session_pointer(self):
        anchor = await self.semantic_pair()
        session = storage.create_revision_session(anchor['id'])
        op, _ = storage.begin_revision_operation(session['id'], 'sync-fail', anchor['id'], 'revise',
            {'feedback': '改为16:00', 'mode': 'targeted', 'artifact_id': self.original['id']})
        with patch.object(service, 'evolve_contract', return_value=self.contract), \
             patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '16:00'}), \
             patch.object(service, '_rewrite', side_effect=HTTPException(502, 'companion failed')):
            await revision_jobs.run(op['id'])
        self.assertEqual(storage.get_revision_operation(op['id'])['status'], 'failed')
        self.assertEqual(storage.get_revision_session(session['id'])['artifact_id'], anchor['id'])
        self.assertIn('14:00', Document(forge.safe_deliverable_path(self.original['file_name'])).paragraphs[0].text)

    async def test_multi_add_is_one_snapshot_with_independent_dependencies(self):
        with patch.object(forge, 'forge_file', side_effect=self.generate_file), \
             patch.object(service, '_validate', side_effect=self.accepted):
            result = await outputs.edit(self.original, '配套PPT和HTML', self.contract,
                action='add', target_kinds=['pptx', 'html'], operation_id='multi-add')
        records = outputs.members(result['artifact'])
        self.assertEqual(len(records), 3)
        self.assertTrue(all(r['sync_from'] == 'source' for r in records[1:]))
        self.assertFalse(result['artifact']['regen']['bundle_pending'])

    def large_word(self):
        path = Path(forge.safe_deliverable_path(self.original['file_name']))
        doc = Document()
        for i in range(30):
            run = doc.add_paragraph().add_run(f'第{i}段 ' + '保留正文' * 800)
            run.bold = True
        doc.paragraphs[0].add_run('BEGIN-MARK')
        doc.paragraphs[-1].add_run('END-MARK')
        doc.save(path)
        return path

    def chunk_reply(self, system, user, **kwargs):
        data = json.loads(user)
        if 'manifest' in data:
            return {'notes': '只修改首尾标记，保留所有正文与样式', 'new_files': [], 'delete_files': []}
        edits = []
        for unit in data['units']:
            for marker in ('BEGIN-MARK', 'END-MARK'):
                if marker in unit['text']:
                    edits.append({'unit_id': unit['id'], 'old_text': marker, 'replacement': marker + '-CHANGED'})
        return {'edits': edits, 'summary': '已核对本块'}

    async def test_large_word_covers_last_chunk_and_preserves_original_and_styles(self):
        path = self.large_word(); original = path.read_bytes()
        with patch.object(llm, 'chat_json', side_effect=self.chunk_reply) as model:
            fields = await service._rewrite(self.original, self.contract, '修改首尾标记')
        coverage = fields['regen']['chunk_edit']
        self.assertGreaterEqual(coverage['total_chunks'], 10)
        self.assertEqual(coverage['reviewed_chunks'], list(range(1, coverage['total_chunks'] + 1)))
        self.assertEqual(model.call_count, coverage['total_chunks'] + 1)
        doc = Document(forge.safe_deliverable_path(fields['file_name']))
        self.assertIn('BEGIN-MARK-CHANGED', doc.paragraphs[0].text)
        self.assertIn('END-MARK-CHANGED', doc.paragraphs[-1].text)
        self.assertEqual(doc.paragraphs[15].text, Document(path).paragraphs[15].text)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertEqual(path.read_bytes(), original)

    def test_late_chunk_failure_does_not_merge_anything(self):
        path = self.large_word(); original = path.read_bytes()
        def reply(system, user, **kwargs):
            data = json.loads(user)
            if data.get('chunk') == 5:
                return {'invalid': 'response'}
            return self.chunk_reply(system, user, **kwargs)
        with patch.object(llm, 'chat_json', side_effect=reply), self.assertRaises(HTTPException):
            chunks.edit(self.original, '修改标记', self.contract)
        self.assertEqual(path.read_bytes(), original)
        self.assertFalse(list(Path(forge.DELIVERABLES_DIR).rglob('分块修订.docx')))

    def test_chunk_rejects_cross_chunk_unit(self):
        with patch.object(llm, 'chat_json', side_effect=[{'notes': 'plan'},
            {'edits': [{'unit_id': 'paragraph:999@0', 'old_text': '14:00', 'replacement': '16:00'}]}]):
            with self.assertRaises(HTTPException):
                chunks.edit(self.original, '修改时间', self.contract)

    def test_chunk_excel_format_and_value_preserve_formula(self):
        name = 'chunk.xlsx'; wb = Workbook(); sheet = wb.active
        sheet['A1'] = 10; sheet['B1'] = '=A1*2'; wb.save(forge.output_path(name)); wb.close()
        item = artifacts.add_artifact('final', '表格', 3, '', file_name=name, sprint_id=self.sid)
        def reply(system, user, **kwargs):
            data = json.loads(user)
            if 'manifest' in data: return {'notes': '改A1数值及格式'}
            unit = next(u for u in data['units'] if u['key'].endswith('!A1'))
            return {'edits': [{'unit_id': unit['id'], 'old_text': '10', 'replacement': '20'},
                {'unit_id': unit['id'], 'operation': 'format', 'format': {'bold': True, 'fill': 'FFDD00'}}]}
        with patch.object(llm, 'chat_json', side_effect=reply):
            result = chunks.edit(item, '修改数值并加粗填色', self.contract)
        wb = load_workbook(forge.safe_deliverable_path(result['file_name']))
        try:
            self.assertEqual(wb.active['A1'].value, 20)
            self.assertTrue(wb.active['A1'].font.bold)
            self.assertEqual(wb.active['A1'].fill.fgColor.rgb[-6:], 'FFDD00')
            self.assertEqual(wb.active['B1'].value, '=A1*2')
        finally: wb.close()

    def images(self):
        directory = Path(self.temp.name) / 'images'; directory.mkdir(exist_ok=True)
        before, after = directory / 'before.png', directory / 'after.png'
        Image.new('RGB', (300, 400), 'white').save(before)
        image = Image.new('RGB', (300, 400), 'white')
        ImageDraw.Draw(image).rectangle((40, 50, 100, 90), fill='black'); image.save(after)
        return directory, before, after

    def test_pixel_regions_and_added_removed_pages(self):
        directory, before, after = self.images()
        same = visual.page_regions(before, before, directory / 'same', 1)
        changed = visual.page_regions(before, after, directory / 'changed', 1)
        self.assertEqual(same['regions'], [])
        self.assertEqual(changed['status'], 'changed')
        self.assertTrue(any(x0 <= 40 and y0 <= 50 and x1 >= 101 and y1 >= 91 for x0, y0, x1, y1 in changed['regions']))
        self.assertGreater(changed['changed_pixels'], 0)
        self.assertTrue((directory / 'changed' / changed['files']['after']).is_file())
        self.assertEqual(visual.page_regions(None, after, directory / 'add', 2)['status'], 'added')
        self.assertEqual(visual.page_regions(before, None, directory / 'remove', 2)['status'], 'removed')

    def test_render_cache_binds_current_ids_and_safe_page_path(self):
        _, before, after = self.images()
        other = artifacts.add_artifact('final', '其他版本', 3, '', file_name=self.original['file_name'], sprint_id=self.sid)
        with patch.object(visual, '_pages', side_effect=[[before], [after]]):
            result = visual.compare(other['id'], self.original['id'])
        self.assertEqual(result['status'], 'ready', result)
        self.assertEqual(result['changed_pages'], 1)
        self.assertTrue(visual.page_path(other['id'], self.original['id'], result['render_id'], 1, 'after').is_file())
        with self.assertRaises(HTTPException):
            visual.page_path(other['id'], self.original['id'], '../escape', 1, 'after')
        with patch.object(visual, '_pages', side_effect=RuntimeError('Office renderer unavailable')):
            missing = visual.compare(other['id'], self.original['id'])
        self.assertEqual(missing['status'], 'unavailable')
        self.assertIn('unavailable', missing['reason'])

    def test_source_change_is_rejected_before_chunk_merge(self):
        source = chunks.Source(self.original)
        Path(source.path).write_bytes(b'changed while reading')
        with self.assertRaises(HTTPException) as error:
            chunks._merge(source, [], {}, [], self.contract)
        self.assertEqual(error.exception.status_code, 409)

    def code_source(self):
        root = Path(self.temp.name) / 'project'; root.mkdir()
        (root / 'app.py').write_text('value = 1\n', encoding='utf-8')
        (root / 'test_app.py').write_text('assert 1 == 1\n', encoding='utf-8')
        (root / 'obsolete.txt').write_text('obsolete', encoding='utf-8')
        (root / 'asset.bin').write_bytes(b'\x00\xff\x01')
        name = 'chunk-code.zip'
        with zipfile.ZipFile(forge.output_path(name), 'w') as archive:
            for path in root.iterdir(): archive.write(path, path.name)
        item = artifacts.add_artifact('final', '代码', 3, '', file_name=name,
            sprint_id=self.sid, regen={'project_dir': str(root), 'runtime': 'python'})
        return item, root

    def test_chunk_code_multifile_merge_preserves_binary_and_tests(self):
        item, root = self.code_source()
        def reply(system, user, **kwargs):
            data = json.loads(user)
            if 'manifest' in data:
                return {'notes': '修改常量，新增说明，删除废弃文件', 'runtime': 'python',
                    'new_files': [{'path': 'docs/info.txt', 'content': 'new documentation'}], 'delete_files': ['obsolete.txt']}
            unit = next(u for u in data['units'] if u['key'] == 'file:app.py')
            return {'edits': [{'unit_id': unit['id'], 'old_text': 'value = 1', 'replacement': 'value = 2'}]}
        with patch.object(llm, 'chat_json', side_effect=reply), \
             patch.object(project_agent, 'evaluate', return_value={'status': 'passed', 'errors': []}) as evaluate:
            result = chunks.edit(item, '修改项目', self.contract)
        evaluate.assert_called_once()
        with zipfile.ZipFile(forge.safe_deliverable_path(result['file_name'])) as archive:
            self.assertEqual(archive.read('app.py'), (root / 'app.py').read_bytes().replace(b'value = 1', b'value = 2'))
            self.assertEqual(archive.read('test_app.py'), (root / 'test_app.py').read_bytes())
            self.assertEqual(archive.read('asset.bin'), b'\x00\xff\x01')
            self.assertEqual(archive.read('docs/info.txt'), b'new documentation')
            self.assertNotIn('obsolete.txt', archive.namelist())
        self.assertEqual((root / 'app.py').read_text(encoding='utf-8'), 'value = 1\n')

    def test_chunk_code_rejects_test_weakening_and_live_directory_drift(self):
        item, root = self.code_source()
        def reply(system, user, **kwargs):
            data = json.loads(user)
            if 'manifest' in data: return {'notes': 'remove tests', 'delete_files': ['test_app.py']}
            return {'edits': []}
        with patch.object(llm, 'chat_json', side_effect=reply), self.assertRaises(HTTPException):
            chunks.edit(item, '修改代码', self.contract)
        source = chunks.Source(item)
        (root / 'app.py').write_text('value = 99\n', encoding='utf-8')
        with self.assertRaises(HTTPException) as error:
            chunks._merge(source, [], {}, [], self.contract)
        self.assertEqual(error.exception.status_code, 409)

    def test_large_source_summary_reduces_context_after_reading_all_chunks(self):
        self.large_word()
        reads = []
        def reply(system, user, **kwargs):
            if user.startswith('{'):
                data = json.loads(user); reads.append(data['chunk'])
                return {'facts': f'块{data["chunk"]}来源记录 ' + '信息' * 4500}
            return {'facts': '已合并所有来源标记与数据'}
        with patch.object(llm, 'chat_json', side_effect=reply):
            result = chunks.summarize(self.original)
        self.assertGreaterEqual(len(reads), 10)
        self.assertEqual(reads, list(range(1, len(reads) + 1)))
        self.assertLess(len(result), 60000)
        self.assertIn('非原始全文', result)

    async def test_semantic_dependencies_cascade_in_order(self):
        anchor = await self.semantic_pair()
        with patch.object(forge, 'forge_file', side_effect=self.generate_file), \
             patch.object(service, '_validate', side_effect=self.accepted):
            anchor = (await outputs.edit(anchor, '新增HTML摘要', self.contract,
                action='add', target_kind='html', operation_id='chain'))['artifact']
        seen = []
        async def synchronize(sid, old, source, contract, operation_id, version):
            seen.append((service.kind_of(old), source['id']))
            return artifacts.add_artifact('final', '同步', 3, '', file_name=old['file_name'],
                sprint_id=sid, parent_id=old['id'], version=version, regen={'bundle_pending': True})
        with patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '16:00'}), \
             patch.object(outputs, 'sync_content', side_effect=synchronize), \
             patch.object(service, '_validate', side_effect=self.accepted):
            result = await outputs.edit(anchor, '改时间', self.contract, artifact_id=self.original['id'],
                operation_id='cascade', mode='targeted')
        current = {o['kind']: o['artifact_id'] for o in result['outputs']}
        self.assertEqual(seen, [('pptx', current['docx']), ('html', current['pptx'])])


if __name__ == '__main__':
    unittest.main()
