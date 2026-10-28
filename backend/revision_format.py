"""Native Office formatting on a copy. No arbitrary code or OOXML from models."""
import copy
import re
from pathlib import Path

from fastapi import HTTPException


def _number(value, low, high, label):
    if type(value) not in (int, float) or not low <= value <= high:
        raise ValueError(f'{label}必须在{low}到{high}之间')
    return value


def _color(value):
    if not isinstance(value, str) or not re.fullmatch(r'[0-9a-fA-F]{6}', value):
        raise ValueError('颜色需要六位十六进制值')
    return value.upper()


def _styles(patch):
    styles = {key: patch[key] for key in ('color', 'fill', 'font_size', 'font_name', 'bold', 'italic',
              'alignment', 'number_format', 'line_spacing', 'column_width', 'row_height') if key in patch}
    if not styles:
        raise ValueError('缺少具体格式要求')
    for key in ('color', 'fill'):
        if key in styles:
            styles[key] = _color(styles[key])
    if 'font_size' in styles:
        _number(styles['font_size'], 5, 96, '字号')
    for key in ('bold', 'italic'):
        if key in styles and type(styles[key]) is not bool:
            raise ValueError('加粗/斜体必须为布尔值')
    if 'font_name' in styles and (not isinstance(styles['font_name'], str) or not 1 <= len(styles['font_name']) <= 100):
        raise ValueError('字体名称无效')
    if 'alignment' in styles and styles['alignment'] not in ('left', 'center', 'right', 'justify'):
        raise ValueError('对齐方式无效')
    if 'line_spacing' in styles:
        _number(styles['line_spacing'], .5, 3, '行距')
    if 'column_width' in styles:
        _number(styles['column_width'], 1, 100, '列宽')
    if 'row_height' in styles:
        _number(styles['row_height'], 5, 200, '行高')
    if 'number_format' in styles and (not isinstance(styles['number_format'], str) or len(styles['number_format']) > 200):
        raise ValueError('数字格式无效')
    return styles


def _font(font, style, *, excel=False, presentation=False):
    if 'font_size' in style:
        if excel:
            font.sz = style['font_size']
        else:
            from docx.shared import Pt
            font.size = Pt(style['font_size'])
    if 'font_name' in style:
        font.name = style['font_name']
    if 'bold' in style:
        font.bold = style['bold']
    if 'italic' in style:
        font.italic = style['italic']
    if 'color' in style:
        if excel:
            font.color = style['color']
        else:
            if presentation:
                from pptx.dml.color import RGBColor
            else:
                from docx.shared import RGBColor
            font.color.rgb = RGBColor.from_string(style['color'])


def apply(kind: str, source: Path, target: Path, patch: dict) -> dict:
    operations = patch.get('formats', [patch])
    if not isinstance(operations, list) or not 1 <= len(operations) <= 100:
        raise ValueError('格式操作数量无效')
    if kind == 'docx':
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        document = Document(source)
        for operation in operations:
            style = _styles(operation)
            if set(style) - {'color', 'font_size', 'font_name', 'bold', 'italic', 'alignment', 'line_spacing'}:
                raise ValueError('Word 不支持该格式字段')
            if operation.get('scope') == 'document':
                paragraphs = list(document.paragraphs)
                paragraphs += [p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
            elif 'table' in operation:
                if any(type(operation.get(key)) is not int or operation[key] < 1 for key in ('table', 'row', 'column')):
                    raise ValueError('表格、行、列序号必须从1开始')
                table = document.tables[int(operation['table']) - 1]
                paragraphs = list(table.cell(int(operation['row']) - 1, int(operation['column']) - 1).paragraphs)
            else:
                number = operation.get('paragraph')
                if type(number) is not int or not 1 <= number <= len(document.paragraphs):
                    raise ValueError('需要有效的顶层段落序号')
                paragraphs = [document.paragraphs[number - 1]]
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    _font(run.font, style)
                if 'alignment' in style:
                    paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                                           'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[style['alignment']]
                if 'line_spacing' in style:
                    paragraph.paragraph_format.line_spacing = style['line_spacing']
        document.save(target)
        from candidate_revision import _sync_word_source
        _sync_word_source(target)
    elif kind == 'pptx':
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        deck = Presentation(source)
        for operation in operations:
            style = _styles(operation)
            if set(style) - {'color', 'fill', 'font_size', 'font_name', 'bold', 'italic', 'alignment'}:
                raise ValueError('PPT 不支持该格式字段')
            page = operation.get('page')
            if type(page) is not int or not 1 <= page <= len(deck.slides):
                raise ValueError('需要有效的PPT页码')
            slide = deck.slides[page - 1]
            index = operation.get('shape')
            if index is not None:
                if type(index) is not int or not 1 <= index <= len(slide.shapes):
                    raise ValueError('PPT形状序号无效')
                shapes = [slide.shapes[index - 1]]
            else:
                old = str(operation.get('old_text') or '')
                shapes = [shape for shape in slide.shapes if shape.has_text_frame and old and old in shape.text]
                if len(shapes) != 1:
                    raise ValueError('PPT格式操作需要唯一文本或明确形状序号')
            for shape in shapes:
                if 'fill' in style:
                    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(style['fill'])
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            _font(run.font, style, presentation=True)
                        if 'alignment' in style:
                            paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                                'right': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}[style['alignment']]
        deck.save(target)
    elif kind == 'xlsx':
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        workbook = load_workbook(source)
        try:
            for operation in operations:
                style = _styles(operation)
                if 'line_spacing' in style:
                    raise ValueError('Excel不支持段落行距')
                location = str(operation.get('location') or '')
                if '!' not in location:
                    raise ValueError('Excel格式操作需要工作表!单元格或区域')
                sheet, area = location.rsplit('!', 1)
                if sheet not in workbook or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?', area):
                    raise ValueError('Excel区域无效')
                left, top, right, bottom = range_boundaries(area)
                if right > 16384 or bottom > 1048576 or (right - left + 1) * (bottom - top + 1) > 5000:
                    raise ValueError('格式区域超过5000个单元格或超出工作表边界')
                worksheet = workbook[sheet]
                for row in worksheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                    for cell in row:
                        font = copy.copy(cell.font); _font(font, style, excel=True); cell.font = font
                        if 'fill' in style:
                            cell.fill = PatternFill('solid', fgColor=style['fill'])
                        if 'number_format' in style:
                            cell.number_format = style['number_format']
                        if 'alignment' in style:
                            alignment = copy.copy(cell.alignment)
                            alignment.horizontal = style['alignment']; cell.alignment = alignment
                if 'column_width' in style:
                    for column in range(left, right + 1):
                        worksheet.column_dimensions[get_column_letter(column)].width = style['column_width']
                if 'row_height' in style:
                    for row in range(top, bottom + 1):
                        worksheet.row_dimensions[row].height = style['row_height']
            workbook.save(target)
        finally:
            workbook.close()
    else:
        raise HTTPException(422, '该类型没有可编辑的原生Office格式，请修改其源文件')
    return {'operation': 'format', 'location': patch.get('location') or '原生格式', 'formats': operations}
