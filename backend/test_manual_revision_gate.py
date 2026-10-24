"""Manual edits keep a version chain and invalidate collaborative approval."""

import asyncio
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("OFFICE_DB_PATH", str(Path(tempfile.mkdtemp()) / "test.db"))

from docx import Document

import artifacts
import artifact_quality
import forge
import main
import storage
import handoffs
import review_decision
import verification


class ManualRevisionGateTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.patches = [
            patch.object(storage, "DB_PATH", str(root / "office.db")),
            patch.object(forge, "WORKSPACE_DIR", str(root / "workspace")),
            patch.object(forge, "DELIVERABLES_DIR", str(root / "deliverables")),
            patch.object(artifacts, "emit"),
        ]
        for item in self.patches:
            item.start()
        storage.init()
        self.sid = storage.create_sprint("联调通知 Word")
        storage.update_sprint(self.sid, execution_style="collaborative")
        storage.save_contract(self.sid, {"status": "ready", "raw_requirement": "联调通知",
                                         "requirements": [{"requirement_id": "R1", "text": "联调通知",
                                                           "status": "ready", "source": {"type": "user"}}]})
        artifacts.set_sprint(self.sid)
        source = Path(forge.output_path("notice.docx"))
        source.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_paragraph("联调时间 14:00")
        document.add_paragraph("其他安排保持不变")
        document.save(source)
        self.original = artifacts.add_artifact(
            "final", "联调通知", 6, "已发布", file_name="notice.docx",
            quality_status="passed", regen={"mode": "published_candidate", "candidate_id": "old"},
            sprint_id=self.sid,
        )

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def test_located_revision_runs_qa_lead_and_preserves_old_file(self):
        response = {"location": "第1段", "old_text": "14:00", "replacement": "15:00"}
        qa_reports = []
        lead_reviews = []

        def fake_verify(candidate, contract, digest, evidence):
            qa_reports.append((candidate["id"], digest, evidence.get("candidate_sha256")))
            report = {"status": "passed", "candidate_id": candidate["id"],
                      "candidate_sha256": digest, "items": [], "checks": [], "reason": ""}
            path = forge.safe_deliverable_path(candidate["file_name"])
            directory = artifact_quality.report_directory(path, forge.WORKSPACE_DIR)
            directory.mkdir(parents=True, exist_ok=True)
            (directory / "qa-report.json").write_text(json.dumps(report), encoding="utf-8")
            return report

        def fake_lead(candidate, contract, report, records, digest, companions=None, evidence=None):
            lead_reviews.append((candidate["id"], digest, report["candidate_sha256"]))
            decision = {"status": "passed", "candidate_id": candidate["id"],
                        "candidate_sha256": digest, "qa_candidate_sha256": digest,
                        "companion_hashes": [], "issues": [], "reason": "同版本复核通过"}
            path = forge.safe_deliverable_path(candidate["file_name"])
            directory = artifact_quality.report_directory(path, forge.WORKSPACE_DIR)
            (directory / "lead-decision.json").write_text(json.dumps(decision), encoding="utf-8")
            return decision

        with (patch.object(main, "chat_json", return_value=response),
              patch.object(forge, "extract_preview", return_value="联调时间 14:00"),
              patch.object(verification, "verify", side_effect=fake_verify),
              patch.object(review_decision, "decide", side_effect=fake_lead)):
            result = asyncio.run(main.revise_artifact(
                self.original["id"], main.RevisionRequest(instruction="把联调时间改为 15:00")
            ))
        revised = result["artifact"]
        self.assertEqual(revised["quality_status"], "passed")
        self.assertEqual(revised["parent_id"], self.original["id"])
        self.assertEqual(revised["regen"]["mode"], "manual_revision")
        self.assertNotEqual(revised["regen"]["sha256"], "")
        self.assertEqual([p.text for p in Document(forge.safe_deliverable_path(revised["file_name"])).paragraphs],
                         ["联调时间 15:00", "其他安排保持不变"])
        self.assertEqual(Document(forge.safe_deliverable_path(self.original["file_name"])).paragraphs[0].text,
                         "联调时间 14:00")
        self.assertEqual(len(qa_reports), 1)
        self.assertEqual(len(lead_reviews), 1)
        self.assertEqual(qa_reports[0][1], lead_reviews[0][1])
        provenance = asyncio.run(main.artifact_provenance(revised["id"]))
        self.assertEqual(provenance["candidate"]["id"], revised["id"])
        self.assertEqual(provenance["qa"]["status"], "passed")
        self.assertEqual(provenance["lead"]["status"], "passed")

    def test_structural_pass_cannot_approve_unreviewed_artifact(self):
        with patch("artifact_quality.load_report", return_value={"status": "passed"}):
            item = artifacts.add_artifact("final", "待复核", 6, "内容", file_name=self.original["file_name"],
                                          quality_status="unchecked", sprint_id=self.sid)
        self.assertEqual(item["quality_status"], "unchecked")

    def test_manual_model_call_is_blocked_when_sprint_budget_is_exhausted(self):
        storage.update_sprint(self.sid, budget=0.25, usage_cost=0.25)
        with patch.object(main.llm, "usage_total_cost", return_value=1.0):
            token = main._manual_call_guard(self.sid)
            try:
                with self.assertRaises(main.BudgetExceeded):
                    main.llm._call_guard.get()()
            finally:
                main.llm.reset_call_guard(token)

    def test_manual_model_call_stops_once_remaining_budget_is_consumed(self):
        storage.update_sprint(self.sid, budget=0.25, usage_cost=0.15)
        with patch.object(main.llm, "usage_total_cost", side_effect=[1.0, 1.0, 1.10]):
            token = main._manual_call_guard(self.sid)
            try:
                guard = main.llm._call_guard.get()
                guard()
                with self.assertRaises(main.BudgetExceeded):
                    guard()
            finally:
                main.llm.reset_call_guard(token)

    def test_failed_final_still_shows_qa_and_lead_candidate(self):
        path = forge.safe_deliverable_path(self.original["file_name"])
        digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        candidate = artifacts.add_artifact(
            "candidate", "候选", 6, "候选", file_name=self.original["file_name"],
            regen={"mode": "candidate", "file_kind": "docx", "sha256": digest}, sprint_id=self.sid,
        )
        directory = artifact_quality.report_directory(path, forge.WORKSPACE_DIR)
        directory.mkdir(parents=True, exist_ok=True)
        (directory / "qa-report.json").write_text(json.dumps({"status": "unchecked", "candidate_sha256": digest}))
        (directory / "lead-decision.json").write_text(json.dumps({"status": "failed", "candidate_sha256": digest}))
        final = artifacts.add_artifact("final", "交付待修正", 6, "未通过", degraded=True,
                                       degraded_reason="任务未通过", quality_status="failed", sprint_id=self.sid)
        result = asyncio.run(main.artifact_provenance(final["id"]))
        self.assertEqual(result["candidate"]["id"], candidate["id"])
        self.assertEqual(result["qa"]["status"], "unchecked")
        self.assertEqual(result["lead"]["status"], "failed")
        self.assertEqual(result["failure_reason"], "任务未通过")


if __name__ == "__main__":
    unittest.main()
