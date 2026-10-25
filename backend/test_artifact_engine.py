"""多格式成品引擎的真实文件与闭环回归；不调用外部模型。"""
import asyncio
from contextlib import closing
import json
import os
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("OFFICE_DB_PATH", str(Path(tempfile.mkdtemp()) / "test.db"))

import artifact_quality
import forge
import main
import office_runtime
from pdf_builder import build_pdf
import project_agent
from spreadsheet_builder import safe_formula


class ArtifactEngineTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.patch = patch.object(forge, "WORKSPACE_DIR", str(self.root / "workspace"))
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def test_unrequested_blank_fields_are_removed_from_notification(self):
        content = ("# 接口联调通知\n\n11 月 4 日联调。\n\n"
                   "会议链接：待补充\n报名：请填写\n日期：TBD\n"
                   "发出方：待确认\n"
                   "| 联调地点 | 线上会议室 A（接入链接/平台：待确认） |\n"
                   "| 前端 | 提交接口清单（提交渠道：待确认） |\n"
                   "## 三、需说明的信息\n入会方式待确认。\n上述信息确认后同步。\n")
        cleaned = forge._strip_unrequested_docx_placeholders(content, "生成一页接口联调通知")
        self.assertNotIn("待补充", cleaned)
        self.assertNotIn("请填写", cleaned)
        self.assertNotIn("TBD", cleaned)
        self.assertNotIn("待确认", cleaned)
        self.assertIn("线上会议室 A", cleaned)
        self.assertIn("提交接口清单", cleaned)
        self.assertNotIn("需说明的信息", cleaned)
        self.assertNotIn("上述信息确认后同步", cleaned)
        self.assertIn("11 月 4 日联调", cleaned)
        collaborative_spec = "需求：生成接口联调通知。\n团队各环节的产出：\n发出方：待确认"
        self.assertNotIn("待确认", forge._strip_unrequested_docx_placeholders(content, collaborative_spec))
        self.assertEqual(forge._strip_unrequested_docx_placeholders(content, "制作待填写的通知模板"), content)

    def test_strip_uses_raw_requirement_in_collaborative_assembly_spec(self):
        # graph.py 交付节点的真实 spec：前缀含“团队各环节”，团队产出正文也含“待确认”，
        # 两者都不应让短通知的严格清理失效（W1 Sprint S1792901859326 回归）。
        spec = (
            "根据需求与团队各环节的实际产出，汇编最终成品。需求：生成一份给研发与测试同事的 Word 通知，"
            "说明 11 月 4 日 14:00 至 16:00 进行接口联调，地点为线上会议室 A。控制在一页，不要封面。\n"
            "任务完成情况：T1 事实核对（completed）。\n\n"
            "团队各环节的产出（请吸收其中的实质内容）：\n"
            "【后端 · T1 · 来源 ['用户原话']】\n未知信息保持待确认。\n\n"
            "重要：必须把实际内容完整做出来，而不是描述怎么做或给模板。"
        )
        content = ("# 接口联调通知\n\n联调时间：11 月 4 日 14:00 至 16:00\n\n"
                   "## 待确认事项\n- 线上会议室 A 的会议链接与入会方式：待确认。\n"
                   "- 通知发出方与落款日期：待确认。\n")
        cleaned = forge._strip_unrequested_docx_placeholders(content, spec)
        self.assertNotIn("待确认", cleaned)
        self.assertNotIn("待确认事项", cleaned)
        self.assertIn("11 月 4 日 14:00 至 16:00", cleaned)
        upload_spec = (
            "用户上传了附件素材，要求基于附件加工改造产出新成品。\n附件内容：\n需求：以附件为准\n\n"
            "加工需求：生成接口联调通知，说明联调安排。\n"
            "团队各环节的产出（请吸收其中的实质内容）：\n未知信息保持待确认。"
        )
        self.assertNotIn("待确认", forge._strip_unrequested_docx_placeholders(content, upload_spec))

    def test_word_review_claims_must_match_actual_table_and_workflow_scope(self):
        path = self.root / "通知.docx"
        forge._build_docx(
            "# 接口联调通知\n\n请准时参会。\n\n"
            "| 角色 | 职责与截止时间 |\n|---|---|\n"
            "| 前端 | 提交清单 |\n| 后端 | 准备环境 |\n| 测试 | 汇总问题 |",
            str(path),
        )
        report = {"kind": "docx", "requirement": "生成通知。请由团队协作完成。"}
        self.assertTrue(artifact_quality._review_issue_contradicted_by_word_file(
            {"problem": "正文缺失角色-职责对照表，未整合表格"}, report, path))
        self.assertTrue(artifact_quality._review_issue_contradicted_by_word_file(
            {"problem": "正文未体现协作痕迹及交接"}, report, path))
        self.assertFalse(artifact_quality._review_issue_contradicted_by_word_file(
            {"problem": "日期与用户要求不一致"}, report, path))

    def test_pdf_is_searchable_and_all_pages_render(self):
        import pymupdf
        path = self.root / "报告.pdf"
        text = "# 验收报告\n\n库存不足时应拒绝借阅，不得重复扣减库存。\n\n## 测试结果\n\n已执行基础校验，并发测试仍待执行。"
        self.assertTrue(build_pdf(forge._parse_docx_markdown(text), str(path)))
        report = artifact_quality.inspect(path, "pdf", "", self.root)
        self.assertEqual(report["status"], "passed")
        with pymupdf.open(path) as document:
            self.assertIn("库存不足", "".join(p.get_text() for p in document))
            self.assertTrue(document.get_toc())
        self.assertEqual(report["metrics"]["pages"], len(report["pages"]))

    def test_word_uses_native_headings_without_forced_cover_or_title_border(self):
        from docx import Document
        from docx.oxml.ns import qn
        path = self.root / "通知.docx"
        forge._build_docx("# 联调安排\n\n请于周五完成接口联调。\n\n## 验收条件\n\n接口错误能在页面显示明确原因。", str(path))
        document = Document(path)
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertTrue(any(p.style.name == "Heading 1" for p in document.paragraphs))
        self.assertFalse(list(document.styles["Title"]._element.iter(qn("w:pBdr"))))
        self.assertFalse(document._element.xpath("//w:br[@w:type='page']"))

    def test_word_and_pdf_remove_model_handoff_tail_before_rendering(self):
        from docx import Document
        draft = "# 我的校园生活\n\n清晨的校园很安静。\n\n（全文约 800 字）\n\n说明：以上为作文正文，可直接复制到 Word 中使用。\n\n---"
        path = self.root / "作文.docx"
        with patch.object(forge, "_gen_docx", return_value=draft):
            self.assertTrue(forge._forge_once("docx", "写800字作文", str(path), None))
        body = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("清晨的校园", body)
        self.assertNotIn("全文约", body)
        self.assertNotIn("可直接复制", body)
        self.assertNotIn("全文约", (artifact_quality.report_directory(path, forge.WORKSPACE_DIR) / "source.md").read_text(encoding="utf-8"))
        pdf_path = self.root / "作文.pdf"
        with patch.object(forge, "_gen_docx", return_value=draft):
            self.assertTrue(forge._forge_once("pdf", "写800字作文", str(pdf_path), None))
        import pymupdf
        with pymupdf.open(pdf_path) as document:
            pdf_text = "\n".join(page.get_text() for page in document)
        self.assertIn("清晨的校园", pdf_text)
        self.assertNotIn("全文约", pdf_text)
        self.assertNotIn("可直接复制", pdf_text)

    def test_quality_inspection_flags_handoff_in_an_existing_word_file(self):
        path = self.root / "旧作文.docx"
        forge._build_docx("# 我的校园生活\n\n放学后我走过操场。\n\n（全文约 800 字）", str(path))
        with patch.object(office_runtime, "convert", side_effect=RuntimeError("测试中跳过页面渲染")):
            report = artifact_quality.inspect(path, "docx", "写作文", self.root)
        self.assertFalse(next(c for c in report["checks"] if c["label"] == "交付正文纯净度")["ok"])

    def test_slides_and_sheets_reject_model_handoff_text(self):
        slides = {"title": "方案", "slides": [{"heading": "结论", "bullets": ["以上为示例内容，可直接复制到 Word 中使用。"]}]}
        sheets = [{"name": "汇总", "headers": ["项目", "值"], "rows": [["说明：以上为作文正文，可直接复制到 Word 中使用。", 1]]}]
        with patch.object(forge, "_gen_pptx", return_value=slides):
            self.assertFalse(forge._forge_once("pptx", "方案", str(self.root / "方案.pptx"), None))
        with patch.object(forge, "_gen_xlsx", return_value=sheets):
            self.assertFalse(forge._forge_once("xlsx", "统计", str(self.root / "统计.xlsx"), None))

    def test_pptx_generation_passes_traceable_cheap_fallback_to_json_model(self):
        meta = {}
        payload = {"title": "方案", "slides": [{"heading": "结论", "bullets": ["先完成状态通知"]}]}
        with (patch.object(artifact_quality, "contract", return_value={"unit": "页", "count": 1}),
              patch.object(forge, "chat_json", return_value=payload) as request):
            result = forge._gen_pptx("生成 1 页方案", 6, generation_meta=meta)
        self.assertEqual(result["slides"][0]["heading"], "结论")
        self.assertIs(request.call_args.kwargs["meta"], meta)
        self.assertEqual(request.call_args.kwargs["fallback_roster"], 1)

    def test_slide_renderer_does_not_add_branding_or_unrequested_closing_slide(self):
        from pptx import Presentation
        path = self.root / "方案.pptx"
        payload = {"title": "方案", "include_cover": False, "slides": [
            {"heading": "问题", "bullets": ["库存不足导致借阅失败。"], "table": None, "note": ""},
            {"heading": "处理", "bullets": ["先校验库存，再确认借阅。"], "table": None, "note": ""},
        ]}
        self.assertTrue(forge._build_pptx(payload, str(path)))
        deck = Presentation(path)
        self.assertEqual(len(deck.slides), 2)
        text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertNotIn("自动生成", text)
        self.assertNotIn("谢谢观看", text)

    def test_ppt_cover_is_included_in_footer_page_numbers(self):
        from pptx import Presentation
        path = self.root / "封面页码.pptx"
        payload = {"title": "方案", "include_cover": True, "slides": [
            {"heading": "结论", "bullets": ["先完成状态通知"], "table": None, "note": ""},
        ]}
        self.assertTrue(forge._build_pptx(payload, str(path)))
        deck = Presentation(path)
        self.assertEqual(len(deck.slides), 2)
        content_text = "\n".join(shape.text for shape in deck.slides[1].shapes if shape.has_text_frame)
        self.assertIn("第 2 页 / 共 2 页", content_text)

    def test_pptx_quality_rejects_wrong_footer_total_and_non_native_chart(self):
        payload = {"title": "三周汇报", "include_cover": True, "slides": [
            {"heading": "趋势", "bullets": ["缺陷持续处理"], "table": None, "note": ""},
        ]}
        path = self.root / "错误页脚.pptx"
        self.assertTrue(forge._build_pptx(payload, str(path)))
        from pptx import Presentation
        deck = Presentation(path)
        footer = next(shape for shape in deck.slides[1].shapes
                      if shape.has_text_frame and "共 2 页" in shape.text)
        footer.text = "第 2 页 / 共 1 页"
        deck.save(path)
        with patch.object(office_runtime, "convert", side_effect=RuntimeError("renderer unavailable")):
            report = artifact_quality.inspect(path, "pptx", "制作 2 页含封面的 PPT，包含原生图表", self.root)
        checks = {item["label"]: item for item in report["checks"]}
        self.assertFalse(checks["页脚总页数一致"]["ok"])
        self.assertFalse(checks["需求中的原生图表"]["ok"])

    def test_document_quality_rejects_a_date_missing_from_the_request(self):
        path = self.root / "通知.docx"
        forge._build_docx("# 接口联调通知\n\n联调时间：11 月 4 日。\n\n研发团队 / 2024 年 11 月", str(path))
        with patch.object(office_runtime, "convert", side_effect=RuntimeError("renderer unavailable")):
            report = artifact_quality.inspect(path, "docx", "写通知，时间为 11 月 4 日", self.root)
        check = next(item for item in report["checks"] if item["label"] == "日期有来源依据")
        self.assertFalse(check["ok"], check)
        self.assertIn("2024年11月", check["detail"])

    def test_quality_inspection_invalidates_old_schema_reports(self):
        path = self.root / "缓存迁移.pdf"
        path.write_bytes(b"candidate")
        previous = {"schema": 2, "requirement": "same", "capabilities": office_runtime.capabilities(),
                    "pages": []}
        fresh = {"schema": 6, "status": "failed", "checks": []}
        with patch.object(artifact_quality, "load_report", return_value=previous), \
             patch.object(artifact_quality, "_inspect", return_value=fresh) as inspect:
            report = artifact_quality.inspect(path, "pdf", "same", self.root)
        inspect.assert_called_once()
        self.assertEqual(report["schema"], 6)

    def test_code_review_flags_handoff_in_readme(self):
        (self.root / "app.py").write_text("print('ok')", encoding="utf-8")
        (self.root / "README.md").write_text("以上为示例代码，可直接复制到 Word 中使用。", encoding="utf-8")
        checks = project_agent.static_check(self.root)
        self.assertFalse(next(c for c in checks if c["label"] == "交付内容纯净度")["ok"])

    def test_ordered_lists_do_not_receive_a_second_bullet_and_headings_do_not_start_at_level_three(self):
        from docx import Document
        blocks = forge._parse_docx_markdown("# 通知\n\n### 结论\n\n内容。\n\n## 安排\n\n1. 完成鉴权\n2. 开始联调")
        self.assertEqual(next(b["type"] for b in blocks if b.get("text") == "结论"), "h2")
        ordered = [b for b in blocks if b["type"] == "ordered"]
        self.assertEqual([b["text"] for b in ordered], ["完成鉴权", "开始联调"])
        path = self.root / "编号.docx"
        forge._build_docx("# 通知\n\n1. 完成鉴权\n2. 开始联调", str(path))
        doc = Document(path)
        self.assertTrue(all(p.style.name == "List Number" for p in doc.paragraphs if p.text in ("完成鉴权", "开始联调")))

    @unittest.skipUnless(office_runtime.find_office(), "需安装 Office 渲染与重算引擎")
    def test_real_formulas_recalculate_and_charts_are_native(self):
        from openpyxl import load_workbook
        path = self.root / "预算.xlsx"
        data = [{"name": "预算", "headers": ["项目", "预算", "实际", "差异"],
                 "rows": [["开发", 1000, 800, "=B2-C2"], ["测试", 600, 650, "=B3-C3"],
                          ["合计", "=SUM(B2:B3)", "=SUM(C2:C3)", "=B4-C4"]],
                 "column_types": {"预算": "currency", "实际": "currency", "差异": "currency"},
                 "charts": [{"type": "bar", "title": "预算比较", "data_columns": [2, 3], "category_column": 1}]}]
        forge._build_xlsx(data, str(path))
        with closing(load_workbook(path, data_only=True)) as wb:
            self.assertEqual(wb["预算"]["D2"].value, 200)
            self.assertEqual(wb["预算"]["D3"].value, -50)
            self.assertEqual(wb["预算"]["D4"].value, 150)
        report = artifact_quality.inspect(path, "xlsx", "包含公式和图表的预算表", self.root)
        self.assertEqual(report["status"], "passed", report["checks"])
        self.assertEqual(report["metrics"]["formulas"], 5)
        self.assertEqual(report["metrics"]["charts"], 1)

    def test_formula_network_and_external_references_are_rejected(self):
        for formula in ('=WEBSERVICE("https://example.test")', "='[secret.xlsx]Sheet1'!A1", "=cmd|' /C calc'!A0"):
            with self.assertRaises(ValueError):
                safe_formula(formula)
        self.assertEqual(safe_formula("=SUM(B2:B4)"), "=SUM(B2:B4)")

    def test_xlsx_schema_preserves_formulas_formats_and_charts(self):
        payload = {"sheets": [{"name": "统计", "headers": ["数量", "金额"], "rows": [[1, "=A2*10"], [2, "=A3*10"]],
                              "formats": {"金额": "0.00"}, "charts": [{"type": "bar", "data_columns": [2]}]}]}
        sanitized = forge._sanitize_xlsx(payload)
        self.assertEqual(sanitized[0]["formats"], payload["sheets"][0]["formats"])
        self.assertEqual(sanitized[0]["charts"], payload["sheets"][0]["charts"])
        self.assertEqual(sanitized[0]["rows"][0][1], "=A2*10")

    def test_unavailable_renderer_is_recorded_as_unchecked(self):
        path = self.root / "文档.docx"
        forge._build_docx("# 简报\n\n今天完成接口开发，明天开始联调。", str(path))
        with patch.object(office_runtime, "convert", side_effect=RuntimeError("渲染引擎不可用")):
            report = artifact_quality.inspect(path, "docx", "", self.root)
        self.assertEqual(report["status"], "unchecked")
        self.assertFalse(report["ok"])

    def test_explicit_requirements_are_not_replaced_by_fixed_minimum_size(self):
        self.assertEqual(artifact_quality.contract("写约200字的通知", "docx")["count"], 200)
        self.assertEqual(artifact_quality.contract("制作6页PPT", "pptx")["count"], 6)
        self.assertEqual(artifact_quality.contract("生成5页含封面的PPT", "pptx")["count_scope"], "total")
        self.assertEqual(artifact_quality.contract("生成5页内容并加封面的PPT", "pptx")["count_scope"], "content")
        self.assertNotIn("不少于 1500", forge._docx_content_prompt("一段简短通知"))

    def test_content_feedback_causes_a_revised_real_document(self):
        drafts = []
        def generate(kind, spec, path, roster, feedback=""):
            drafts.append(feedback)
            return forge._build_docx("# 风险说明\n\n" + ("库存竞争风险通过事务与并发测试处理。" if feedback else "本轮完成开发。"), path)
        def review(report, *args):
            if len(drafts) == 1:
                report.update(ok=False, status="failed", review={"status": "failed", "issues": [{"location": "正文", "problem": "缺少风险说明", "fix": "补充库存竞争与处理办法"}]})
            return report
        with patch.object(forge, "_forge_once", side_effect=generate), patch.object(artifact_quality, "model_review", side_effect=review):
            name = forge.forge_file("docx", "补充风险说明", "风险.docx")
        self.assertEqual(len(drafts), 2)
        self.assertIn("缺少风险说明", drafts[1])
        self.assertIn("库存竞争", forge.extract_preview(name))
        self.assertEqual(len(forge.self_check(name)["attempts"]), 2)

    def test_pdf_and_code_requests_have_real_routes(self):
        import graph
        self.assertEqual(main._detect_file_kind("生成一份PDF报告"), "pdf")
        self.assertEqual(graph.detect_final_kind("导出PDF"), "pdf")
        self.assertEqual(main._detect_file_kind("开发一个Python脚本"), "code")

    def test_project_paths_reject_traversal_and_reserved_devices(self):
        for name in ("../outside.py", "C:/secret.py", "/tmp/run.py", "src/../../run.py", "CON.py", "src/NUL.txt"):
            with self.assertRaises(ValueError):
                project_agent.safe_path(self.root, name)

    def test_static_checks_do_not_execute_generated_python(self):
        marker = self.root / "executed.txt"
        (self.root / "app.py").write_text(f"from pathlib import Path\nPath({str(marker)!r}).write_text('ran')\n")
        checks = project_agent.static_check(self.root)
        self.assertFalse(marker.exists())
        self.assertTrue(any(c["label"] == "自动化测试文件" and c["ok"] is False for c in checks))

    def test_missing_sandbox_cannot_be_reported_as_runtime_pass(self):
        (self.root / "app.py").write_text("def add(a,b):\n    return a+b\n")
        (self.root / "README.md").write_text("python -m unittest discover -s tests")
        (self.root / "tests").mkdir()
        (self.root / "tests/test_app.py").write_text("from app import add\ndef test_add():\n    assert add(1,2)==3\n")
        with patch.object(project_agent, "docker_available", return_value=False):
            report = project_agent.evaluate(self.root, [], "python")
        self.assertEqual(report["status"], "unchecked")
        self.assertFalse(report["ok"])

    def test_project_generation_repairs_actual_syntax_error(self):
        project = self.root / "project"
        team = self.root / "team"
        team.mkdir()
        files = ["app.py", "README.md", "tests/test_app.py"]
        def response(*args, **kwargs):
            self.calls += 1
            app = "def add(a,b)\n    return a+b" if self.calls == 1 else "def add(a,b):\n    return a+b"
            return f"===FILE: app.py\n{app}\n===END\n===FILE: README.md\nRun tests with unittest.\n===END\n===FILE: tests/test_app.py\nfrom app import add\ndef test_add():\n    assert add(1,2)==3\n===END"
        self.calls = 0
        with patch.object(project_agent, "chat_json", return_value={"runtime": "python", "files": [{"path": name} for name in files]}), patch.object(project_agent, "chat_text", side_effect=response), patch.object(project_agent, "docker_available", return_value=False):
            chosen, report = project_agent.assemble("实现加法函数并测试", team, project)
        self.assertEqual(self.calls, 2)
        self.assertNotEqual(chosen, project)
        self.assertEqual(report["status"], "unchecked")
        self.assertTrue(next(c for c in report["checks"] if c["label"] == "语法与解析")["ok"])

    def test_code_runtime_uses_a_network_disabled_test_container_and_no_host_environment(self):
        commands = []
        def run(command, **kwargs):
            commands.append(command)
            if hasattr(kwargs.get("stdout"), "write"):
                kwargs["stdout"].write(b"Ran 1 test in 0.1s\nOK\n")
            return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")
        with patch.object(project_agent, "docker_available", return_value=True), patch.object(project_agent.shutil, "which", return_value="docker"), patch.object(project_agent.subprocess, "run", side_effect=run):
            report = project_agent.sandbox_check(self.root, "python")
        self.assertTrue(report["ok"])
        stages = [command for command in commands if "run" in command]
        self.assertIn("--network=none", stages[1])
        self.assertIn("--read-only", stages[1])
        self.assertIn("--user=65534:65534", stages[1])
        self.assertFalse(any("API_KEY" in part or "docker.sock" in part for command in stages for part in command))

    def test_quality_report_and_real_page_api(self):
        import httpx
        name = forge._unique_name("验收.pdf")
        path = forge.output_path(name)
        build_pdf(forge._parse_docx_markdown("# 验收说明\n\n基础测试通过，并发场景仍待验证。"), path)
        artifact_quality.inspect(path, "pdf", "", forge.WORKSPACE_DIR)
        async def request():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
                with patch.object(main, "_find_artifact", return_value={"id": "TEST", "file_name": name}):
                    result = await client.get("/api/artifacts/TEST/quality")
                    page = await client.get("/api/artifacts/TEST/quality/pages/1")
                    missing = await client.get("/api/artifacts/TEST/quality/pages/999")
            return result, page, missing
        result, page, missing = asyncio.run(request())
        self.assertEqual(result.json()["status"], "passed")
        self.assertEqual(page.headers["content-type"], "image/png")
        self.assertTrue(page.content.startswith(b"\x89PNG"))
        self.assertEqual(missing.status_code, 404)

    def test_model_call_guard_stops_network_requests(self):
        import llm
        def stopped():
            raise RuntimeError("已达到预算或迭代已停止")
        token = llm.set_call_guard(stopped)
        try:
            with patch.object(llm, "_resolve", return_value=("https://example.test/v1", "test-key", "test-model")), patch.object(llm, "_resolve_backup", return_value=None), patch.object(llm.httpx, "post") as post:
                result = llm.chat_text("测试", [{"role": "user", "content": "测试"}], fallback="已停止")
            self.assertEqual(result, "已停止")
            post.assert_not_called()
        finally:
            llm.reset_call_guard(token)


if __name__ == "__main__":
    unittest.main()
