"""Revision regression tests: temporary DB/files, mocked models, no Sprint run."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document

os.environ.setdefault('OFFICE_DB_PATH', str(Path(tempfile.mkdtemp(prefix='revision-tests-')) / 'test.db'))

import artifact_quality
import artifacts
import candidate_revision
import forge
import main
import revision_jobs
import revision_service as service
import runtime_state
import storage
import review_decision


class RevisionSessionTest(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.patches = [patch.object(storage, 'DB_PATH', str(root / 'test.db')),
                        patch.object(forge, 'DELIVERABLES_DIR', str(root / 'deliverables')),
                        patch.object(forge, 'WORKSPACE_DIR', str(root / 'workspace')),
                        patch.object(artifacts, 'emit'), patch.object(runtime_state, 'running', False)]
        for item in self.patches:
            item.start()
        storage.init()
        self.sid = storage.create_sprint('生成 Word 联调通知，约1000字')
        storage.update_sprint(self.sid, execution_style='collaborative')
        self.contract = {'status': 'ready', 'raw_requirement': '约1000字；其他安排保持不变',
                         'attachment_context': '📎《数据.txt》第1行：\n联调时间14:00',
                         'requirements': [
                             {'requirement_id': 'R1', 'text': '约1000字', 'status': 'ready',
                              'source': {'type': 'user'}, 'acceptance_method': '检查字数'},
                             {'requirement_id': 'R2', 'text': '其他安排保持不变', 'status': 'ready',
                              'source': {'type': 'user'}, 'acceptance_method': '检查内容'},
                         ]}
        storage.save_contract(self.sid, self.contract)
        path = Path(forge.output_path('notice.docx'))
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_paragraph('联调时间14:00')
        doc.add_paragraph('其他安排保持不变')
        doc.save(path)
        self.original = artifacts.add_artifact('final', '通知', 6, '摘要', file_name='notice.docx',
            quality_status='passed', regen={'mode': 'published_candidate'}, sprint_id=self.sid)
        self.session = storage.create_revision_session(self.original['id'])

    async def asyncTearDown(self):
        await revision_jobs.shutdown()
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def queue(self, request='request-1', feedback='修改时间'):
        return storage.begin_revision_operation(self.session['id'], request, self.original['id'],
                                                'revise', {'feedback': feedback, 'mode': 'targeted'})[0]

    def test_session_restore_and_request_deduplication(self):
        self.assertEqual(storage.create_revision_session(self.original['id'])['id'], self.session['id'])
        first = self.queue()
        repeated, created = storage.begin_revision_operation(self.session['id'], 'request-1', self.original['id'],
            'revise', {'feedback': '修改时间', 'mode': 'targeted'})
        self.assertFalse(created)
        self.assertEqual(first['id'], repeated['id'])
        with self.assertRaises(ValueError):
            self.queue(feedback='不同请求内容')

    def test_conflicts_do_not_create_another_job(self):
        self.queue()
        with self.assertRaises(ValueError):
            self.queue(request='request-2')
        self.assertEqual(len(storage.revision_operations(self.session['id'])), 1)

    def test_failed_job_keeps_current_and_accepted_versions(self):
        op = self.queue()
        storage.finish_revision_operation(op['id'], error='模型不可用')
        session = storage.get_revision_session(self.session['id'])
        self.assertEqual(session['artifact_id'], self.original['id'])
        self.assertEqual(session['accepted_artifact_id'], self.original['id'])
        self.assertEqual(session['history'][0]['status'], 'failed')

    def test_failed_candidate_advances_edit_pointer_but_not_accepted_pointer(self):
        op = self.queue()
        candidate = {**self.original, 'id': 'candidate', 'version': 2, 'quality_status': 'failed'}
        storage.finish_revision_operation(op['id'], artifact=candidate, summary='生成但未通过')
        session = storage.get_revision_session(self.session['id'])
        self.assertEqual(session['artifact_id'], 'candidate')
        self.assertEqual(session['accepted_artifact_id'], self.original['id'])
        with self.assertRaises(ValueError):
            self.queue(request='stale-tab')

    def test_updated_contract_preserves_sources_and_retires_old_length(self):
        plan = {'target_kind': 'docx', 'audience': '团队', 'purpose': '通知',
                'requirements': [{'source': 'user', 'source_quote': '压缩到500字',
                                  'acceptance_method': '统计正文为500字'}],
                'missing_info': [], 'allowed_assumptions': [], 'superseded_ids': ['R1']}
        with patch.object(main, 'chat_json', return_value=plan):
            contract = service.evolve_contract(self.original, '压缩到500字')
        self.assertNotIn('1000字', contract['raw_requirement'])
        self.assertIn('500字', contract['raw_requirement'])
        self.assertIn('其他安排保持不变', contract['raw_requirement'])
        self.assertEqual(contract['attachment_context'], self.contract['attachment_context'])
        self.assertEqual(storage.get_sprint(self.sid)['contract'], self.contract)

    async def test_three_local_edits_use_latest_bytes_and_preserve_original(self):
        current = self.original
        with patch.object(service, '_validate'):
            for index, (old, new) in enumerate([('14:00', '15:00'), ('15:00', '16:00'), ('16:00', '17:00')], 2):
                with patch.object(main, 'chat_json', return_value={
                    'operation': 'replace', 'old_text': old, 'replacement': new, 'location': '第1段'}):
                    result = await service.revise(current['id'], f'把时间改为{new}', contract=self.contract)
                current = result['artifact']
                self.assertEqual(current['version'], index)
                self.assertIn(new, Document(forge.safe_deliverable_path(current['file_name'])).paragraphs[0].text)
                source = artifact_quality.report_directory(forge.safe_deliverable_path(current['file_name']), forge.WORKSPACE_DIR) / 'source.md'
                self.assertIn(new, source.read_text(encoding='utf-8'))
                self.assertNotIn(old, source.read_text(encoding='utf-8'))
        self.assertIn('14:00', Document(forge.safe_deliverable_path(self.original['file_name'])).paragraphs[0].text)

    async def test_macro_after_local_rebuilds_a_real_file_from_current_content(self):
        def generate(kind, spec, filename, roster):
            self.assertIn('15:00', spec)
            name = 'rewritten.docx'
            doc = Document(); doc.add_paragraph('联调时间15:00；全文已重写')
            doc.save(forge.output_path(name))
            return name
        with patch.object(service, '_validate'), patch.object(main, 'chat_json', return_value={
                'old_text': '14:00', 'replacement': '15:00'}):
            local = (await service.revise(self.original['id'], '改为15:00', contract=self.contract))['artifact']
        with patch.object(service, '_validate'), patch.object(forge, 'forge_file', side_effect=generate):
            macro = (await service.revise(local['id'], '整体重写', contract=self.contract))['artifact']
        self.assertNotEqual(local['file_name'], macro['file_name'])
        self.assertIn('全文已重写', Document(forge.safe_deliverable_path(macro['file_name'])).paragraphs[0].text)

    def test_word_format_is_applied_to_real_runs_only(self):
        filename, reason, _ = candidate_revision.apply(self.sid, 'docx', self.original['file_name'],
            {'operation': 'format', 'paragraph': 2, 'color': 'FF0000', 'font_size': 16})
        self.assertFalse(reason)
        doc = Document(forge.safe_deliverable_path(filename))
        self.assertEqual(str(doc.paragraphs[1].runs[0].font.color.rgb), 'FF0000')
        self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt, 16)
        self.assertEqual(doc.paragraphs[0].text, '联调时间14:00')
        self.assertIn('color:#FF0000', forge.preview_html(filename))
        self.assertIn('font-size:16.0pt', forge.preview_html(filename))

    def test_qa_pass_without_visual_evidence_does_not_skip_lead(self):
        import hashlib
        path = forge.safe_deliverable_path(self.original['file_name'])
        digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        candidate = {**self.original, 'regen': {'file_kind': 'docx'}}
        with patch.object(artifact_quality, 'load_report', return_value={
                'text': '通知', 'checks': [{'ok': None}], 'pages': []}), \
             patch.object(review_decision.llm, 'chat_json', return_value={
                 'verdict': 'passed', 'note': '模型认为通过', 'issues': []}) as model:
            decision = review_decision.decide(candidate, self.contract,
                {'status': 'passed', 'candidate_sha256': digest}, [], digest)
        self.assertEqual(model.call_count, 1)
        self.assertEqual(decision['status'], 'unchecked')

    async def test_rollback_then_next_edit_uses_restored_content(self):
        with patch.object(service, '_validate'), patch.object(main, 'chat_json', return_value={
                'old_text': '14:00', 'replacement': '15:00'}):
            revised = (await service.revise(self.original['id'], '修改时间', contract=self.contract))['artifact']
            restored = (await service.rollback(revised['id'], self.original['id']))['artifact']
            with patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '18:00'}):
                final = (await service.revise(restored['id'], '修改时间', contract=self.contract))['artifact']
        self.assertEqual(final['version'], 4)
        self.assertIn('18:00', Document(forge.safe_deliverable_path(final['file_name'])).paragraphs[0].text)

    async def test_background_completion_recovers_session_and_reuses_result(self):
        op = self.queue()
        with patch.object(service, 'evolve_contract', return_value=self.contract), patch.object(service, '_validate'), \
             patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '15:00'}):
            await revision_jobs.run(op['id'])
            await revision_jobs.run(op['id'])
        snapshot = revision_jobs.session_snapshot(self.session['id'])
        self.assertEqual(len(snapshot['history']), 1)
        self.assertNotEqual(snapshot['artifact_id'], self.original['id'])
        self.assertEqual(snapshot['operations'][0]['status'], 'completed')

    def test_recovery_adopts_existing_candidate_without_calling_model(self):
        op = self.queue()
        candidate = artifacts.add_artifact('final', '中断候选', 6, '内容', file_name=self.original['file_name'],
            parent_id=self.original['id'], version=2, sprint_id=self.sid,
            regen={'revision_operation_id': op['id']})
        revision_jobs.recover_interrupted()
        session = storage.get_revision_session(self.session['id'])
        self.assertEqual(session['artifact_id'], candidate['id'])
        self.assertEqual(session['accepted_artifact_id'], self.original['id'])
        self.assertEqual(session['history'][0]['status'], 'interrupted')
        self.assertFalse(storage.revision_operations(active=True))

    async def test_api_validates_requests_and_reports_conflicts_without_model(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://test') as client:
            opened = await client.post('/api/revision/session', json={'artifact_id': self.original['id']})
            self.assertEqual(opened.status_code, 200)
            with patch.object(revision_jobs, 'dispatch') as dispatch:
                request = {'feedback': '修改时间', 'request_id': 'http-1', 'expected_artifact_id': self.original['id']}
                first = await client.post(f'/api/revision/session/{self.session["id"]}/message', json=request)
                retry = await client.post(f'/api/revision/session/{self.session["id"]}/message', json=request)
                self.assertEqual(first.status_code, 202)
                self.assertEqual(first.json(), retry.json())
                self.assertEqual(dispatch.call_count, 1)
                conflict = await client.post(f'/api/revision/session/{self.session["id"]}/message', json={**request, 'request_id': 'http-2'})
                self.assertEqual(conflict.status_code, 409)
            missing = await client.get('/api/revision/session/missing')
            self.assertEqual(missing.status_code, 404)


if __name__ == '__main__':
    unittest.main()
