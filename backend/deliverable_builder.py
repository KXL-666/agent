"""确定性渲染器：Markdown / JSON 负载 → docx / pptx / xlsx 真实文件（不调用 LLM）

只做「结构解析 + 版式渲染」。内容由谁生成、是否修订，由 deliverable_generator 决定；
python-docx / python-pptx / openpyxl 都在函数内延迟导入，缺库时由调用方兜底。
"""

import os
import re

# 统一视觉规范：三类产物共享同一套配色，保证交付物风格一致
COLOR_PRIMARY = "1F4E79"   # 深蓝：封面/页眉带/表头
COLOR_ORANGE = "ED7D31"    # 点缀橙：装饰条
COLOR_LIGHT = "DEEBF7"     # 浅蓝：副标题/隔行
COLOR_GRAY = "808080"      # 辅助灰
FONT_MAIN = "微软雅黑"

# 支持的成品类型（生成入口与自检共用同一份口径）
CONTENT_KINDS = ("docx", "pptx", "xlsx", "pdf", "html")


def _strip_fences(text: str) -> str:
    """剥掉模型偶发输出的 ```markdown 围栏"""
    text = (text or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*\n?", "", text)
        text = re.sub(r"\n?```\s*$", "", text)
    return text.strip()


def _clean_inline(text: str) -> str:
    """去掉行内 Markdown 标记（加粗/行内代码/链接），返回纯文本"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return text.strip()


def _bold_segments(text: str) -> list[tuple[str, bool]]:
    """把含 **加粗** 的文本拆成 [(片段, 是否加粗)]，供渲染器分 run 输出"""
    parts = re.split(r"\*\*(.+?)\*\*", text)
    return [
        (_clean_inline(part), bool(i % 2))
        for i, part in enumerate(parts)
        if part
    ]


def _parse_docx_markdown(text: str) -> list[dict]:
    """把 Markdown 文档解析为块序列：h1/h2/h3/p/bullet/table"""
    blocks: list[dict] = []
    table_buf: list[str] = []
    code_lines: list[str] | None = None

    def flush_table() -> None:
        nonlocal table_buf
        if not table_buf:
            return
        rows: list[list[str]] = []
        for row in table_buf:
            cells = [_clean_inline(c) for c in row.strip().strip("|").split("|")]
            # 跳过分隔行 | --- | --- |
            if all(re.fullmatch(r":?-{2,}:?", c or "---") for c in cells):
                continue
            rows.append(cells)
        if rows:
            blocks.append({"type": "table", "rows": rows})
        table_buf = []

    for raw in _strip_fences(text).replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if line.startswith("```"):
            flush_table()
            if code_lines is None:
                code_lines = []
            else:
                blocks.append({"type": "code", "text": "\n".join(code_lines)})
                code_lines = None
            continue
        if code_lines is not None:
            code_lines.append(raw)
            continue
        if line.startswith("|") and line.endswith("|"):
            table_buf.append(line)
            continue
        flush_table()
        if not line or line in ("---", "***", "___"):
            continue
        if line.startswith("### "):
            blocks.append({"type": "h3", "text": _clean_inline(line[4:])})
        elif line.startswith("## "):
            blocks.append({"type": "h2", "text": _clean_inline(line[3:])})
        elif line.startswith("# "):
            blocks.append({"type": "h1", "text": _clean_inline(line[2:])})
        elif re.match(r"^[-*•]\s+", line):
            blocks.append({"type": "bullet", "text": re.sub(r"^[-*•]\s+", "", line)})
        elif re.match(r"^\d+[.、]\s*\S", line):
            blocks.append({"type": "ordered", "text": re.sub(r"^\d+[.、]\s*", "", line)})
        else:
            blocks.append({"type": "p", "text": line})
    flush_table()
    if code_lines is not None:
        blocks.append({"type": "code", "text": "\n".join(code_lines)})
    seen_major = False
    for block in blocks:
        if block["type"] == "h2":
            seen_major = True
        elif block["type"] == "h3" and not seen_major:
            # 不能在一级章节之前出现其子章节，避免正文从三级标题开始。
            block["type"] = "h2"
    return blocks


def _sanitize_table(table) -> dict | None:
    if not isinstance(table, dict):
        return None
    headers = [str(h).strip()[:30] for h in (table.get("headers") or []) if str(h).strip()][:6]
    if len(headers) < 2:
        return None
    rows = []
    for row in (table.get("rows") or [])[:8]:
        if isinstance(row, list):
            cells = ["" if c is None else str(c).strip()[:40] for c in row][: len(headers)]
            cells += [""] * (len(headers) - len(cells))
            rows.append(cells)
    if not rows:
        return None
    return {"headers": headers, "rows": rows}


def _sanitize_pptx(payload, source_spec: str = "") -> dict | None:
    """清洗 LLM 返回的 PPT JSON 负载；结构严重不符时返回 None"""
    if not isinstance(payload, dict):
        return None
    slides_raw = payload.get("slides")
    if not isinstance(slides_raw, list):
        # 兜底：模型换字段名时，取第一个「dict 列表」字段当页序列
        slides_raw = next(
            (
                v for v in payload.values()
                if isinstance(v, list) and v and isinstance(v[0], dict)
            ),
            [],
        )
    slides: list[dict] = []
    for item in slides_raw[:15]:
        if not isinstance(item, dict):
            continue
        heading = str(item.get("heading") or item.get("title") or "").strip()[:40]
        bullets = [
            str(b).strip()
            for b in (item.get("bullets") or item.get("points") or [])
            if str(b).strip()
        ][:7]
        note = str(item.get("note") or "").strip()[:300]
        layout = str(item.get("layout") or "bullets").strip().lower()
        if layout not in ("bullets", "conclusion", "compare", "process", "timeline", "chart", "table"):
            layout = "bullets"
        chart = item.get("chart") if isinstance(item.get("chart"), dict) else None
        if chart:
            categories = chart.get("categories")
            series = chart.get("series")
            source_refs = item.get("source_refs")
            valid = (isinstance(categories, list) and 2 <= len(categories) <= 12
                     and isinstance(series, list) and 1 <= len(series) <= 4
                     and isinstance(source_refs, list) and bool(source_refs))
            normalized = []
            if valid:
                for entry in series:
                    values = entry.get("values") if isinstance(entry, dict) else None
                    if not isinstance(values, list) or len(values) != len(categories) or any(type(v) not in (int, float) for v in values):
                        valid = False
                        break
                    if source_spec and any(not re.search(rf"(?<!\d){re.escape(str(v))}(?!\d)", source_spec) for v in values):
                        valid = False
                        break
                    normalized.append({"name": str(entry.get("name") or "数值")[:30], "values": values})
            chart = {"categories": [str(c)[:30] for c in categories], "series": normalized,
                     "title": str(chart.get("title") or "")[:60]} if valid else None
        if layout == "chart" and not chart:
            layout = "bullets"
        if heading or bullets:
            slides.append({
                "heading": heading or "（未命名章节）",
                "bullets": bullets,
                "table": _sanitize_table(item.get("table")),
                "note": note,
                "layout": layout,
                "claim": str(item.get("claim") or "").strip()[:160],
                "evidence": str(item.get("evidence") or "").strip()[:300],
                "source_refs": [str(ref).strip()[:100] for ref in (item.get("source_refs") or []) if str(ref).strip()][:8],
                "chart": chart,
            })
    if len(slides) < 1:
        return None
    return {
        "title": str(payload.get("title") or "演示文稿").strip()[:40] or "演示文稿",
        "subtitle": str(payload.get("subtitle") or "").strip()[:60],
        "slides": slides,
        "include_cover": bool(payload.get("include_cover", True)),
    }


def _sanitize_xlsx(payload) -> list[dict] | None:
    """清洗 LLM 返回的 Excel JSON 负载为工作表列表；无有效数据返回 None"""
    if not isinstance(payload, dict):
        return None
    sheets_raw = payload.get("sheets")
    if not isinstance(sheets_raw, list):
        sheets_raw = next(
            (
                v for v in payload.values()
                if isinstance(v, list) and v and isinstance(v[0], dict)
            ),
            [],
        )
    sheets: list[dict] = []
    for item in sheets_raw[:4]:
        if not isinstance(item, dict):
            continue
        headers = [str(h).strip()[:30] for h in (item.get("headers") or []) if str(h).strip()][:40]
        if len(headers) < 2:
            continue
        rows: list[list] = []
        for row in (item.get("rows") or [])[:2000]:
            if not isinstance(row, list):
                continue
            cells = [
                c if isinstance(c, (int, float)) else ("" if c is None else str(c).strip()[:2000])
                for c in row
            ][: len(headers)]
            cells += [""] * (len(headers) - len(cells))
            if any(str(c).strip() for c in cells):
                rows.append(cells)
        if len(rows) >= 2:
            sheets.append({
                "name": str(item.get("name") or f"工作表{len(sheets) + 1}").strip()[:20],
                "headers": headers,
                "rows": rows,
                **{key: item[key] for key in ("formats", "column_types", "formulas", "notes") if isinstance(item.get(key), dict)},
                **{key: item[key] for key in ("charts", "validations") if isinstance(item.get(key), list)},
            })
    return sheets or None


def _docx_set_run(run, size: float = 11, bold: bool = False,
                  color: str | None = None, font: str = FONT_MAIN) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _docx_rich_text(paragraph, text: str, size: float = 11,
                    color: str | None = None) -> None:
    """按 **加粗** 分段输出 run"""
    for segment, bold in _bold_segments(text):
        _docx_set_run(paragraph.add_run(segment), size=size, bold=bold, color=color)


def _docx_shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _docx_page_number(doc) -> None:
    """页脚居中页码（PAGE 域）"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")  # 9pt
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLOR_GRAY)
    rpr.append(sz)
    rpr.append(color)
    run.append(rpr)
    fld.append(run)
    footer_p._p.append(fld)


def _build_docx(content: str, output_path: str, cover: bool = False) -> bool:
    """原生标题与导航、按内容分页的表格和正文；无需通用封面。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    blocks = _parse_docx_markdown(content)
    if not blocks:
        return False
    doc = Document()
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        for border in list(doc.styles[style_name]._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.2)
        section.left_margin = section.right_margin = Cm(2.4)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT_MAIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_MAIN)
        fonts = style._element.get_or_add_rPr().rFonts
        for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + theme), None)
        from docx.shared import RGBColor
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.45
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(8)
    title_used = False
    for block in blocks:
        kind = block["type"]
        if kind in ("h1", "h2", "h3"):
            style = "Title" if kind == "h1" and not title_used else {"h1": "Heading 1", "h2": "Heading 1", "h3": "Heading 2"}[kind]
            heading = doc.add_paragraph(block["text"], style=style)
            for run in heading.runs:
                _docx_set_run(run, size={"Title": 25, "Heading 1": 16, "Heading 2": 13}[style], bold=True, color="000000")
            if cover and style == "Title":
                doc.add_page_break()
            title_used = True
        elif kind == "table":
            rows = block["rows"]
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.autofit = False
            weights = [max(3, min(25, max(len(r[c]) if c < len(r) else 0 for r in rows[:15]))) for c in range(cols)]
            for c, weight in enumerate(weights):
                table.columns[c].width = Cm(16.2 * weight / sum(weights))
            for r, row in enumerate(rows):
                properties = table.rows[r]._tr.get_or_add_trPr()
                if r == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                # 超长记录允许跨页，普通记录保持完整。
                if max(map(len, row), default=0) < 350:
                    properties.append(OxmlElement("w:cantSplit"))
                for c, cell in enumerate(table.rows[r].cells):
                    cell.width = Cm(16.2 * weights[c] / sum(weights))
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    text = row[c] if c < len(row) else ""
                    _docx_rich_text(cell.paragraphs[0], text, size=9.5, color="000000")
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(5)
                        paragraph.paragraph_format.space_before = Pt(5)
                    if r == 0:
                        _docx_shade_cell(cell, "DEEBF7")
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                    elif r % 2 == 0:
                        _docx_shade_cell(cell, "F5F7FA")
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif kind in ("bullet", "ordered"):
            paragraph = doc.add_paragraph(style="List Number" if kind == "ordered" else "List Bullet")
            _docx_rich_text(paragraph, block["text"], size=10.5, color="000000")
        elif kind == "code":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name, run.font.size = "Consolas", Pt(9)
        else:
            paragraph = doc.add_paragraph()
            _docx_rich_text(paragraph, block["text"], size=10.5, color="000000")
    doc.core_properties.title = next((b["text"] for b in blocks if b["type"] == "h1"), "文档")
    doc.core_properties.author = ""
    _docx_page_number(doc)
    doc.save(output_path)
    return os.path.getsize(output_path) > 0


def _ppt_rect(slide, x, y, w, h, fill: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False


def _ppt_text(slide, x, y, w, h, lines, size: float, color: str,
              bold: bool = False, align=None, line_spacing: float = 1.15) -> None:
    """lines: [(text, bold, color)] 或纯 str 列表"""
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(lines):
        text, item_bold, item_color = (
            item if isinstance(item, tuple) else (item, bold, color)
        )
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = text
        run.font.name = FONT_MAIN
        run.font.size = Pt(size)
        run.font.bold = item_bold
        run.font.color.rgb = RGBColor.from_string(item_color)


def _build_pptx(payload: dict, output_path: str) -> bool:
    """把 PPT JSON 负载渲染为统一设计风格的 16:9 演示文稿"""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    width, height = prs.slide_width, prs.slide_height

    if payload.get("include_cover", True):
        # ---------- 标题页 ----------
        slide = prs.slides.add_slide(blank)
        _ppt_rect(slide, 0, 0, width, height, COLOR_PRIMARY)
        _ppt_rect(slide, 0, Emu(int(height * 0.68)), width, Pt(6), COLOR_ORANGE)
        _ppt_text(
            slide, Inches(1), Emu(int(height * 0.30)), width - Inches(2), Inches(1.6),
            [payload["title"]], 40, "FFFFFF", bold=True, align=PP_ALIGN.CENTER,
        )
        if payload.get("subtitle"):
            _ppt_text(
                slide, Inches(1), Emu(int(height * 0.50)), width - Inches(2), Inches(0.8),
                [payload["subtitle"]], 18, COLOR_LIGHT, align=PP_ALIGN.CENTER,
            )

    # ---------- 内容页 ----------
    has_cover = bool(payload.get("include_cover", True))
    total = len(payload["slides"]) + int(has_cover)
    for content_index, item in enumerate(payload["slides"], 1):
        page_number = content_index + int(has_cover)
        slide = prs.slides.add_slide(blank)
        # 顶部标题带 + 点缀条
        _ppt_rect(slide, 0, 0, width, Inches(1.0), COLOR_PRIMARY)
        _ppt_rect(slide, 0, Inches(1.0), Inches(2.4), Pt(4), COLOR_ORANGE)
        _ppt_text(
            slide, Inches(0.6), Inches(0.12), width - Inches(1.2), Inches(0.8),
            [item["heading"]], 26, "FFFFFF", bold=True,
        )
        from presentation_layout import render_body
        render_body(slide, item)
        # 演讲备注
        if item.get("note"):
            slide.notes_slide.notes_text_frame.text = item["note"]
        # 页脚：左侧文档名 + 右侧页码
        _ppt_text(
            slide, Inches(0.5), height - Inches(0.45), Inches(6), Inches(0.35),
            [payload["title"]], 10, COLOR_GRAY,
        )
        _ppt_text(
            slide, width - Inches(2.6), height - Inches(0.45), Inches(2.1), Inches(0.35),
            [f"第 {page_number} 页 / 共 {total} 页"], 10, COLOR_GRAY, align=PP_ALIGN.RIGHT,
        )

    try:
        prs.save(output_path)
        return os.path.getsize(output_path) > 0
    except OSError as error:
        print(f"[Forge] pptx 保存失败：{error}")
        return False


def _build_xlsx(sheets: list[dict], output_path: str) -> bool:
    from spreadsheet_builder import build_workbook
    return build_workbook(sheets, output_path)
