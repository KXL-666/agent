"""P1/P2 slide count, editable native chart and source-number guard."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

import forge


P1 = "生成一份 6 页、不带封面的中文 PPT，没有用户量、节省时间或预算数据，不要编造统计。"
P2 = "根据附件生成 5 页含封面的中文 PPT。第 1 周新增缺陷 12，关闭 8；第 2 周新增 9，关闭 11；第 3 周新增 14，关闭 8。"


class PresentationLayoutTest(unittest.TestCase):
    def test_p1_exact_six_pages_without_fake_chart(self):
        payload = {"title": "校园预约系统", "slides": [
            {"heading": f"结论 {i}", "claim": f"第 {i} 页主张", "layout": layout,
             "bullets": ["状态通知", "管理端筛选导出", "导师确认"]}
            for i, layout in enumerate(("conclusion", "compare", "process", "timeline", "bullets", "table"), 1)
        ]}
        cleaned = forge._sanitize_pptx(payload, P1)
        cleaned["include_cover"] = False
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = str(Path(directory) / "p1.pptx")
            self.assertTrue(forge._build_pptx(cleaned, path))
            deck = Presentation(path)
            self.assertEqual(len(deck.slides), 6)
            self.assertFalse(any(shape.has_chart for slide in deck.slides for shape in slide.shapes))

    def test_p2_five_pages_with_native_chart(self):
        chart = {"title": "每周缺陷", "categories": ["第1周", "第2周", "第3周"],
                 "series": [{"name": "新增", "values": [12, 9, 14]},
                            {"name": "关闭", "values": [8, 11, 8]}]}
        payload = {"title": "缺陷处理进展", "slides": [
            {"heading": "三周对比", "claim": "第三周新增增多", "layout": "chart",
             "source_refs": ["附件第1-3周记录"], "chart": chart, "bullets": []},
            *[{"heading": f"分析 {i}", "layout": "conclusion", "claim": "待确认原因", "bullets": ["未提供原因"]}
              for i in range(2, 5)],
        ]}
        cleaned = forge._sanitize_pptx(payload, P2)
        self.assertIsNotNone(cleaned["slides"][0]["chart"])
        cleaned["include_cover"] = True
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = str(Path(directory) / "p2.pptx")
            self.assertTrue(forge._build_pptx(cleaned, path))
            deck = Presentation(path)
            self.assertEqual(len(deck.slides), 5)
            charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
            self.assertEqual(len(charts), 1)
            self.assertEqual(list(charts[0].series[0].values), [12.0, 9.0, 14.0])
        invented = {**payload, "slides": [{**payload["slides"][0],
                    "chart": {**chart, "series": [{"name": "新增", "values": [12, 9, 999]}]}}]}
        self.assertIsNone(forge._sanitize_pptx(invented, P2)["slides"][0]["chart"])

    def test_explicit_cover_count_sent_to_model(self):
        prompts = []

        def fake_json(system, user, **kwargs):
            prompts.append(user)
            return {"title": "测试", "slides": [{"heading": "结论", "bullets": ["事实"]}]}

        with patch.object(forge, "chat_json", side_effect=fake_json):
            no_cover = forge._gen_pptx(P1, 3)
            with_cover = forge._gen_pptx(P2, 3)
        self.assertFalse(no_cover["include_cover"])
        self.assertTrue(with_cover["include_cover"])
        self.assertIn("只生成 6 个 slides", prompts[0])
        self.assertIn("只生成 4 个 slides", prompts[1])

    def test_long_table_keeps_bottom_evidence_inside_slide(self):
        payload = {"title": "任务清单", "include_cover": False, "slides": [{
            "heading": "安排明细", "layout": "bullets", "claim": "任务安排",
            "bullets": [f"要点 {index}" for index in range(5)],
            "table": {"headers": ["任务", "负责人"],
                      "rows": [[f"工作 {index}", "研发"] for index in range(8)]},
            "evidence": "来源：需求附件第 1 页",
        }]}
        cleaned = forge._sanitize_pptx(payload)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = str(Path(directory) / "table.pptx")
            self.assertTrue(forge._build_pptx(cleaned, path))
            slide = Presentation(path).slides[0]
            tables = [shape for shape in slide.shapes if shape.has_table]
            self.assertEqual(len(tables), 1)
            table_shape = tables[0]
            self.assertLessEqual(table_shape.top + table_shape.height, 6.0 * 914400)
            self.assertEqual(len(table_shape.table.rows), 9)
            visible_bullets = [shape.text for shape in slide.shapes
                               if shape.has_text_frame and shape.text.startswith("• ")]
            self.assertLessEqual(len(visible_bullets), 1)


if __name__ == "__main__":
    unittest.main()
