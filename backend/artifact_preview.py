"""产物预览与自检：文本提取、结构化 HTML 预览、Markdown 重渲染、自检说明。

读取一律走 path_manager 的安全路径；预览失败只影响展示，不会破坏已交付文件。
"""

import html as html_lib
import os
import re
import shutil
import tempfile
from contextlib import closing
from pathlib import Path

from path_manager import _unique_name, _workspace_dir, output_path, safe_deliverable_path
from deliverable_builder import CONTENT_KINDS, _build_docx, _parse_docx_markdown


def render_docx_from_markdown(content: str, filename: str) -> str | None:
    """确定性渲染：把 Markdown 文本渲染为带封面/样式/页码的真实 docx（不调用 LLM）

    用于演示模式与产物站内文本编辑（编辑后确定性重渲染为新版本文件）。"""
    filename = _unique_name(filename)
    file_path = output_path(filename)
    if _build_docx(content, file_path):
        import artifact_quality
        directory = artifact_quality.report_directory(file_path, _workspace_dir())
        directory.mkdir(parents=True, exist_ok=True)
        (directory / "source.md").write_text(content, encoding="utf-8")
        return filename
    print(f"[Forge] Markdown 渲染 docx 失败：{filename}")
    return None


def pdf_from_docx(docx_name: str, filename: str = "同源文档.pdf") -> str | None:
    """Convert the exact Word candidate to searchable PDF without rewriting text."""
    import artifact_quality
    import office_runtime

    source = safe_deliverable_path(docx_name)
    if not source:
        return None
    pdf_name = _unique_name(filename)
    destination = output_path(pdf_name)
    try:
        with tempfile.TemporaryDirectory(prefix="word-pdf-") as directory:
            converted = office_runtime.convert(source, directory, "pdf")
            shutil.copy2(converted, destination)
        source_markdown = artifact_quality.report_directory(source, _workspace_dir()) / "source.md"
        if source_markdown.is_file():
            target_quality = artifact_quality.report_directory(destination, _workspace_dir())
            target_quality.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_markdown, target_quality / "source.md")
        return pdf_name
    except (OSError, RuntimeError, ValueError) as error:
        print(f"[Forge] Word 同源 PDF 转换失败：{error}")
        return None


def render_document_from_markdown(content: str, filename: str) -> str | None:
    if not filename.lower().endswith(".pdf"):
        name = render_docx_from_markdown(content, filename)
    else:
        from pdf_builder import build_pdf
        name = _unique_name(filename)
        if not build_pdf(_parse_docx_markdown(content), output_path(name)):
            return None
    if name:
        import artifact_quality
        directory = artifact_quality.report_directory(output_path(name), _workspace_dir())
        directory.mkdir(parents=True, exist_ok=True)
        (directory / "source.md").write_text(content, encoding="utf-8")
        artifact_quality.inspect(output_path(name), Path(filename).suffix.lstrip("."), "", _workspace_dir())
    return name


def extract_preview(file_name: str, limit: int = 1500) -> str:
    """从生成的真实文件中提取文本，作为产物中心的正文预览"""
    path = safe_deliverable_path(file_name)
    if not path:
        return ""
    if file_name.lower().endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)[:limit]
    try:
        if file_name.endswith(".docx"):
            from docx import Document

            paras = [
                p.text.strip()
                for p in Document(path).paragraphs
                if p.text.strip()
            ]
            for table in Document(path).tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paras.append(" | ".join(cells))
            return "\n".join(paras)[:limit]
        if file_name.endswith(".pptx"):
            from pptx import Presentation

            lines: list[str] = []
            for i, slide in enumerate(Presentation(path).slides, 1):
                lines.append(f"【第 {i} 页】")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = "".join(run.text for run in para.runs).strip()
                            if text:
                                lines.append(text)
            return "\n".join(lines)[:limit]
        if file_name.endswith(".xlsx"):
            from openpyxl import load_workbook

            with closing(load_workbook(path, read_only=True, data_only=True)) as wb:
                lines = []
                for ws in wb.worksheets:
                    lines.append(f"【工作表：{ws.title}】")
                    for row in ws.iter_rows(max_row=10, max_col=6, values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            lines.append(" | ".join(cells))
                    lines.append("")
                return "\n".join(lines)[:limit]
    except Exception as error:  # noqa: BLE001 —— 预览失败不影响文件本身
        print(f"[Forge] 预览提取失败 {file_name}：{error}")
    return ""


def preview_html(file_name: str, max_rows: int = 60) -> str:
    """生成结构化 HTML 预览（Word 标题/段落/表格、PPT 分页、Excel 工作表），
    所有文本均转义，前端可直接渲染。"""
    path = safe_deliverable_path(file_name)
    if not path:
        return ""
    esc = html_lib.escape
    try:
        if file_name.lower().endswith('.pdf'):
            from pypdf import PdfReader
            return ''.join(f'<section><h3>第 {index} 页</h3><pre style="white-space:pre-wrap">{esc(page.extract_text() or "该页无可提取文字")}</pre></section>'
                           for index, page in enumerate(PdfReader(path).pages, 1))
        if file_name.endswith(".docx"):
            from docx import Document

            doc = Document(path)
            parts: list[str] = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                style = (p.style.name or "").lower()
                tag = (
                    "h1" if "heading 1" in style or style == "title"
                    else "h2" if "heading 2" in style
                    else "h3" if "heading 3" in style
                    else "p"
                )
                content = esc(text)
                if ''.join(run.text for run in p.runs) == p.text:
                    spans = []
                    for run in p.runs:
                        styles = []
                        rgb = str(run.font.color.rgb or '')
                        if re.fullmatch(r'[0-9A-Fa-f]{6}', rgb):
                            styles.append('color:#' + rgb)
                        if run.font.size is not None:
                            styles.append(f'font-size:{run.font.size.pt}pt')
                        if run.bold:
                            styles.append('font-weight:bold')
                        spans.append(f'<span style="{";".join(styles)}">{esc(run.text)}</span>')
                    content = ''.join(spans)
                parts.append(f"<{tag}>{content}</{tag}>")
            for table in doc.tables:
                rows = []
                for row in table.rows[:max_rows]:
                    cells = "".join(f"<td>{esc(c.text.strip())}</td>" for c in row.cells)
                    rows.append(f"<tr>{cells}</tr>")
                if rows:
                    parts.append(f"<table>{''.join(rows)}</table>")
            return "".join(parts)
        if file_name.endswith(".pptx"):
            from pptx import Presentation

            parts = []
            for index, slide in enumerate(Presentation(path).slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = "".join(run.text for run in para.runs).strip()
                            if text:
                                texts.append(f"<p>{esc(text)}</p>")
                parts.append(
                    f'<section class="preview-slide"><h3>第 {index} 页</h3>{"".join(texts)}</section>'
                )
            return "".join(parts)
        if file_name.endswith(".xlsx"):
            from openpyxl import load_workbook

            with closing(load_workbook(path, read_only=True, data_only=True)) as wb:
                parts = []
                for ws in wb.worksheets:
                    rows = []
                    for row in ws.iter_rows(max_row=max_rows, max_col=8, values_only=True):
                        cells = "".join(
                            f"<td>{esc('' if c is None else str(c))}</td>" for c in row
                        )
                        rows.append(f"<tr>{cells}</tr>")
                    parts.append(
                        f'<section class="preview-sheet"><h3>{esc(ws.title)}</h3>'
                        f'<table>{"".join(rows)}</table></section>'
                    )
                return "".join(parts)
    except Exception as error:  # noqa: BLE001
        print(f"[Forge] HTML 预览失败 {file_name}：{error}")
    return ""


def self_check(file_name: str, kind: str | None = None) -> dict:
    import artifact_quality
    path = safe_deliverable_path(file_name)
    if not path:
        return {"ok": False, "status": "failed", "checks": [{"label": "文件存在", "ok": False, "detail": "文件不存在"}]}
    existing = artifact_quality.load_report(path, _workspace_dir())
    if existing and existing.get("kind") in CONTENT_KINDS:
        import office_runtime
        if existing.get("capabilities") != office_runtime.capabilities():
            return artifact_quality.inspect(path, existing["kind"], existing.get("requirement", ""), _workspace_dir())
    return existing or artifact_quality.inspect(path, kind or Path(file_name).suffix.lstrip("."), "", _workspace_dir())


def self_check_note(result: dict) -> str:
    """把自检结果格式化为可直接附在产物正文后的 Markdown 小节"""
    lines = ["", "", "## 成品自检", ""]
    for item in result.get("checks", []):
        mark = "✅" if item["ok"] is True else "未验证" if item["ok"] is None else "⚠️"
        lines.append(f"- {mark} {item['label']}：{item['detail']}")
    review = result.get("review", {})
    if review.get("status") in ("passed", "failed"):
        lines.append("- 模型成品评审：" + ("通过" if review["status"] == "passed" else "需修正"))
        lines.extend("- " + str(issue) for issue in review.get("issues", []))
    if not result.get("ok"):
        lines.append("")
        lines.append("> 请查看未通过项或未验证项；缺少工具能力时，重新生成不会替代实际验证。")
    return "\n".join(lines)
