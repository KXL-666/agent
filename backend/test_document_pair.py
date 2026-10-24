"""W1 single-page Word and D1 Word/PDF derived from one content source."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import artifact_quality
import forge
import office_runtime


W1 = ("生成一份给研发与测试同事的 Word 通知，说明 11 月 4 日 14:00 至 16:00 进行接口联调，"
      "地点为线上会议室 A。前端在 11 月 3 日 18:00 前提交调用清单，后端在联调开始前准备测试环境，"
      "测试同事记录阻塞问题并在当日 20:00 前汇总。控制在一页，不要封面。")
W1_TEXT = ("# 接口联调通知\n\n研发与测试同事：\n\n"
           "11 月 4 日 14:00 至 16:00 在「线上会议室 A」进行接口联调。\n\n"
           "- 前端：11 月 3 日 18:00 前提交调用清单。\n"
           "- 后端：联调开始前准备测试环境。\n"
           "- 测试：记录阻塞问题，并在当日 20:00 前汇总。")
D1_TEXT = ("# 预约系统试点进度分析\n\n"
           "## 结论\n\n预约表单、手机号格式校验和记录写入接口已完成；短信提醒和取消预约仍未完成。\n\n"
           "## 测试与问题\n\n已执行 20 条用例，16 条通过，4 条失败，通过率为 16/20=80%。\n\n"
           "## 下一步\n\n计划 10 月 28 日前完成取消预约修复，11 月 2 日前完成短信提醒联调。")


class DocumentPairTest(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.old = forge.WORKSPACE_DIR, forge.DELIVERABLES_DIR
        forge.WORKSPACE_DIR = str(Path(self.scratch.name) / "workspace")
        forge.DELIVERABLES_DIR = str(Path(self.scratch.name) / "deliverables")
        Path(forge.WORKSPACE_DIR).mkdir()
        Path(forge.DELIVERABLES_DIR).mkdir()

    def tearDown(self):
        forge.WORKSPACE_DIR, forge.DELIVERABLES_DIR = self.old
        self.scratch.cleanup()

    def _word(self, name, content):
        path = Path(forge.output_path(f"_process/{name}"))
        self.assertTrue(forge._build_docx(content, str(path), cover=False))
        directory = artifact_quality.report_directory(path, forge.WORKSPACE_DIR)
        directory.mkdir(parents=True, exist_ok=True)
        (directory / "source.md").write_text(content, encoding="utf-8")
        return path

    def test_w1_one_page_no_cover(self):
        path = self._word("w1.docx", W1_TEXT)
        document = Document(path)
        self.assertEqual(document.paragraphs[0].text, "接口联调通知")
        self.assertFalse(any("封面" in p.text for p in document.paragraphs))

    def test_w1_page_count_is_verified_or_explicitly_skipped(self):
        if not office_runtime.find_office():
            self.skipTest("Office 渲染器不可用，无法核验真实页数")
        path = self._word("w1-page-check.docx", W1_TEXT)
        report = artifact_quality.inspect(path, "docx", W1, forge.WORKSPACE_DIR)
        self.assertEqual(report["metrics"]["pages"], 1)
        self.assertTrue(all(phrase in report["text"] for phrase in ("14:00", "16:00", "18:00", "20:00")))

    def test_d1_pdf_uses_exact_word_source(self):
        word = self._word("d1.docx", D1_TEXT)
        if not office_runtime.find_office():
            self.skipTest("Word→PDF 转换器不可用")
        pdf_name = forge.pdf_from_docx("_process/d1.docx", "d1.pdf")
        self.assertIsNotNone(pdf_name)
        pdf = Path(forge.output_path(pdf_name))
        word_source = artifact_quality.report_directory(word, forge.WORKSPACE_DIR) / "source.md"
        pdf_source = artifact_quality.report_directory(pdf, forge.WORKSPACE_DIR) / "source.md"
        self.assertEqual(word_source.read_bytes(), pdf_source.read_bytes())
        rendered = office_runtime.render_pdf(pdf, artifact_quality.report_directory(pdf, forge.WORKSPACE_DIR))
        self.assertTrue(rendered["text"].strip())
        self.assertIn("16/20=80%", rendered["text"])

    def test_word_to_pdf_conversion_failure_is_reported_as_missing(self):
        word = self._word("d1-failed.docx", D1_TEXT)
        with patch.object(office_runtime, "convert", side_effect=RuntimeError("模拟转换器不可用")):
            self.assertIsNone(forge.pdf_from_docx("_process/d1-failed.docx", "failed.pdf"))


if __name__ == "__main__":
    unittest.main()
