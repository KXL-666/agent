"""Native editable PowerPoint body layouts selected by the page's message."""

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


PRIMARY = RGBColor(31, 78, 121)
ACCENT = RGBColor(46, 117, 182)
LIGHT = RGBColor(235, 243, 250)
TEXT = RGBColor(45, 55, 65)


def _text(slide, x, y, w, h, value, size=18, bold=False, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = str(value)
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _card(slide, x, y, w, h, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = ACCENT
    _text(slide, x + .2, y + .15, w - .4, h - .3, text, 17)


def _table(slide, data, top=2.15, bottom=6.0):
    headers = data["headers"]
    rows = data["rows"]
    nrows = len(rows) + 1
    desired_height = min(3.6, max(.8, nrows * .46))
    height = min(desired_height, max(.5, bottom - top))
    table = slide.shapes.add_table(nrows, len(headers), Inches(.8), Inches(top), Inches(11.7), Inches(height)).table
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows, 1):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
                    run.font.color.rgb = TEXT


def render_body(slide, item: dict) -> None:
    layout = item.get("layout", "bullets")
    claim = item.get("claim") or item.get("heading") or ""
    bullets = item.get("bullets") or []
    if layout == "chart" and item.get("chart"):
        chart = item["chart"]
        _text(slide, .85, 1.28, 11.7, .65, claim, 23, True, PRIMARY)
        data = CategoryChartData()
        data.categories = chart["categories"]
        for series in chart["series"]:
            data.add_series(series["name"], series["values"])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2.05),
                                       Inches(11.3), Inches(3.8), data)
        frame.chart.has_legend = len(chart["series"]) > 1
        frame.chart.has_title = False
    elif layout == "table" and item.get("table"):
        _text(slide, .85, 1.28, 11.7, .65, claim, 23, True, PRIMARY)
        _table(slide, item["table"])
    elif layout == "conclusion":
        _text(slide, .9, 1.55, 11.5, 1.5, claim, 30, True, PRIMARY)
        for index, bullet in enumerate(bullets[:3]):
            _card(slide, .9, 3.25 + index * .8, 11.5, .67, bullet)
    elif layout == "compare":
        _text(slide, .85, 1.28, 11.7, .7, claim, 23, True, PRIMARY)
        left = bullets[: max(1, (len(bullets) + 1) // 2)]
        right = bullets[len(left):]
        for index, column in enumerate((left, right)):
            x = .85 + index * 6.0
            _card(slide, x, 2.25, 5.55, 3.85, "\n\n".join(column) or "待确认")
    elif layout in ("process", "timeline"):
        _text(slide, .85, 1.28, 11.7, .7, claim, 23, True, PRIMARY)
        steps = bullets[:5]
        if steps:
            width = min(3.5, 11.5 / len(steps) - .18)
            for index, step in enumerate(steps):
                x = .85 + index * (11.5 / len(steps))
                _card(slide, x, 2.55, width, 2.6, f"{index + 1}. {step}")
    else:
        _text(slide, .85, 1.28, 11.7, .7, claim, 23, True, PRIMARY)
        table = item.get("table")
        bottom = 5.95 if item.get("evidence") else 6.8
        table_top = None
        visible_bullets = bullets[:5]
        if table:
            rows = len(table.get("rows", [])) + 1
            table_height = min(3.6, max(.8, rows * .46))
            room = bottom - 2.25
            max_bullets = max(0, min(5, int((room - table_height - .12) // .68)))
            visible_bullets = visible_bullets[:max_bullets]
            table_top = 2.25 + len(visible_bullets) * .68 + .12
        for index, bullet in enumerate(visible_bullets):
            _text(slide, 1, 2.25 + index * .68, 11.2, .58, "• " + bullet, 19)
        if table:
            _table(slide, table, top=table_top or 2.25, bottom=bottom)
    evidence = item.get("evidence") or ""
    if evidence:
        _text(slide, .85, 6.15, 11.7, .55, "依据：" + evidence, 11, color=ACCENT)
