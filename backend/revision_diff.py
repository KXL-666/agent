"""Full-file content and native-format comparisons, independent of card summaries."""
import difflib
import json
import zipfile
from pathlib import Path

from fastapi import HTTPException

import forge
import storage

MAX_ITEMS = 100000
MAX_BYTES = 8 * 1024 * 1024


def _color(value):
    try:
        return str(value.rgb) if value.type is not None else None
    except (AttributeError, ValueError):
        return str(value)


def _font(font):
    return {'name': font.name, 'size': font.size.pt if font.size else None,
            'bold': font.bold, 'italic': font.italic, 'color': _color(font.color)}


def snapshot(artifact: dict) -> list[dict]:
    filename = artifact.get('file_name') or ''
    path = forge.safe_deliverable_path(filename)
    if not path:
        raise HTTPException(404, '对比所需版本文件已丢失')
    kind = Path(filename).suffix.lower()
    records = []
    def add(location, text, format=None):
        records.append({'location': location, 'text': str(text), 'format': format or {}})
        if len(records) > MAX_ITEMS:
            raise HTTPException(422, '文件超过完整对比的条目上限，请拆分文件')
    if kind == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = Document(path)
        for index, block in enumerate(document.element.body, 1):
            if block.tag.endswith('}p'):
                paragraph = Paragraph(block, document)
                add(f'正文块{index}', paragraph.text, {'style': paragraph.style.name,
                    'alignment': str(paragraph.alignment),
                    'line_spacing': paragraph.paragraph_format.line_spacing,
                    'space_before': paragraph.paragraph_format.space_before,
                    'space_after': paragraph.paragraph_format.space_after,
                    'runs': [dict(text=run.text, **_font(run.font)) for run in paragraph.runs]})
            elif block.tag.endswith('}tbl'):
                table = Table(block, document)
                for row, cells in enumerate(table.rows, 1):
                    for column, cell in enumerate(cells.cells, 1):
                        add(f'正文块{index}/表格/{row},{column}', cell.text,
                            {'paragraphs': [{'alignment': str(p.alignment), 'line_spacing': p.paragraph_format.line_spacing,
                                'runs': [dict(text=r.text, **_font(r.font)) for r in p.runs]} for p in cell.paragraphs]})
        for index, section in enumerate(document.sections, 1):
            add(f'页面设置{index}', '', {key: getattr(section, key) for key in (
                'top_margin', 'bottom_margin', 'left_margin', 'right_margin', 'page_width', 'page_height')})
    elif kind == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=False)
        try:
            for sheet in workbook:
                if sheet.max_row * sheet.max_column > MAX_ITEMS:
                    raise HTTPException(422, '工作表超过完整对比的单元格上限')
                for row in sheet:
                    for cell in row:
                        if cell.value is not None or cell.has_style:
                            add(f'{sheet.title}!{cell.coordinate}', cell.value if cell.value is not None else '',
                                {'font': {'name': cell.font.name, 'size': cell.font.sz, 'bold': cell.font.b,
                                          'italic': cell.font.i, 'color': str(cell.font.color)},
                                 'fill': str(cell.fill.fgColor), 'number_format': cell.number_format,
                                 'alignment': str(cell.alignment), 'border': str(cell.border)})
                add(f'{sheet.title}/工作表设置', '', {
                    'columns': {key: value.width for key, value in sheet.column_dimensions.items()},
                    'rows': {str(key): value.height for key, value in sheet.row_dimensions.items()},
                    'merged': str(sheet.merged_cells),
                    'validations': str(sheet.data_validations),
                    'charts': [str(chart) for chart in sheet._charts]})
        finally:
            workbook.close()
    elif kind == '.pptx':
        from pptx import Presentation
        deck = Presentation(path)
        for page, slide in enumerate(deck.slides, 1):
            for index, shape in enumerate(slide.shapes, 1):
                content = shape.text if shape.has_text_frame else ''
                format = {'x': shape.left, 'y': shape.top, 'width': shape.width, 'height': shape.height,
                          'fill_type': str(shape.fill.type) if hasattr(shape, 'fill') else None}
                if shape.has_text_frame:
                    format['paragraphs'] = [{'alignment': str(p.alignment), 'runs': [dict(text=r.text, **_font(r.font)) for r in p.runs]} for p in shape.text_frame.paragraphs]
                if shape.has_table:
                    content += json.dumps([[c.text for c in row.cells] for row in shape.table.rows], ensure_ascii=False)
                if shape.has_chart:
                    content += json.dumps([{'name': s.name, 'values': list(s.values)} for s in shape.chart.series], ensure_ascii=False)
                if hasattr(shape, 'fill') and str(shape.fill.type) == 'SOLID (1)':
                    format['fill'] = _color(shape.fill.fore_color)
                add(f'第{page}页/形状{index}', content, format)
    elif kind == '.pdf':
        from pypdf import PdfReader
        for index, page in enumerate(PdfReader(path).pages, 1):
            add(f'第{index}页', page.extract_text() or '', {'width': float(page.mediabox.width), 'height': float(page.mediabox.height)})
    elif kind == '.zip':
        with zipfile.ZipFile(path) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
            if sum(entry.file_size for entry in entries) > MAX_BYTES:
                raise HTTPException(422, '代码包超过完整对比的大小上限')
            for entry in sorted(entries, key=lambda item: item.filename):
                data = archive.read(entry)
                if b'\x00' in data:
                    import hashlib
                    add(entry.filename, '二进制 SHA-256:' + hashlib.sha256(data).hexdigest())
                else:
                    add(entry.filename, data.decode('utf-8', errors='replace'))
    else:
        if Path(path).stat().st_size > MAX_BYTES:
            raise HTTPException(422, '文件超过完整对比的大小上限')
        add('全文', Path(path).read_text(encoding='utf-8', errors='replace'))
    if len(json.dumps(records, ensure_ascii=False, default=str).encode('utf-8')) > MAX_BYTES:
        raise HTTPException(422, '文件超过完整对比的大小上限')
    return records


def _format_only(value):
    """Run text is retained in evidence but is content, not a style change."""
    if isinstance(value, dict):
        return {key: _format_only(item) for key, item in value.items() if key != 'text'}
    if isinstance(value, list):
        return [_format_only(item) for item in value]
    return value


def compare(artifact_id: str, base_id: str) -> dict:
    current, base = storage.get_artifact(artifact_id), storage.get_artifact(base_id)
    if not current or not base:
        raise HTTPException(404, '版本不存在')
    if current['sprint_id'] != base['sprint_id']:
        raise HTTPException(400, '不能比较不同任务的文件')
    before, after = snapshot(base), snapshot(current)
    left, right = {r['location']: r for r in before}, {r['location']: r for r in after}
    changes = []
    for location in dict.fromkeys([*left, *right]):
        old, new = left.get(location), right.get(location)
        if old == new:
            continue
        changes.append({'location': location, 'type': 'add' if old is None else 'delete' if new is None else 'change',
                        'before': old, 'after': new,
                        'content_changed': old is None or new is None or old['text'] != new['text'],
                        'format_changed': old is not None and new is not None
                            and _format_only(old['format']) != _format_only(new['format'])})
    left_text = '\n'.join(r['location'] + '\n' + r['text'] for r in before)
    right_text = '\n'.join(r['location'] + '\n' + r['text'] for r in after)
    return {'artifact_id': artifact_id, 'base_id': base_id, 'complete': True,
            'changes': changes, 'content_changes': sum(c['content_changed'] for c in changes),
            'format_changes': sum(c['format_changed'] for c in changes),
            'text_diff': '\n'.join(difflib.unified_diff(left_text.splitlines(), right_text.splitlines(),
                                     fromfile=f'v{base.get("version", 1)}', tofile=f'v{current.get("version", 1)}', lineterm='')),
            'visual_note': '此处比较完整文本及可读取的原生格式；实际页面外观请并排预览。'}
