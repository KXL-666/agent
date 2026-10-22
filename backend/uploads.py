"""附件管理：用户上传文件/文件夹，AI 基于附件内容加工改造产出新文件。

原始文件保存在 backend/uploads/，只读使用，永不被 AI 修改。
元数据存于 uploads/index.json。
"""

import json
import os
import time

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
INDEX_PATH = os.path.join(UPLOADS_DIR, "index.json")

os.makedirs(UPLOADS_DIR, exist_ok=True)

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB
TEXT_EXTS = {".txt", ".md", ".py", ".js", ".ts", ".jsx", ".tsx", ".json",
             ".html", ".css", ".java", ".c", ".cpp", ".go", ".rs", ".sql",
             ".yml", ".yaml", ".xml", ".csv", ".sh", ".bat", ".ini", ".log"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff"}


def _load_index() -> list[dict]:
    try:
        with open(INDEX_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return []


def _save_index(items: list[dict]) -> None:
    tmp = INDEX_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    os.replace(tmp, INDEX_PATH)


def _sanitize_name(filename: str) -> str:
    """去掉路径分隔符防穿越，只保留文件名部分"""
    name = filename.replace("\\", "/").split("/")[-1].strip()
    return name or f"unnamed_{int(time.time())}"


def save_upload(filename: str, data: bytes) -> dict | None:
    """保存上传文件并登记元数据；超限返回 None"""
    if len(data) > MAX_FILE_SIZE:
        return None
    name = _sanitize_name(filename)
    upload_id = f"U{int(time.time() * 1000)}"
    disk_name = f"{upload_id}_{name}"
    with open(os.path.join(UPLOADS_DIR, disk_name), "wb") as f:
        f.write(data)
    item = {
        "id": upload_id,
        "name": name,
        "disk_name": disk_name,
        "size": len(data),
        "uploaded_at": time.time(),
    }
    items = _load_index()
    items.append(item)
    _save_index(items)
    return item


def list_uploads() -> list[dict]:
    return [
        {k: v for k, v in item.items() if k != "disk_name"}
        for item in _load_index()
    ]


def delete_upload(upload_id: str) -> bool:
    items = _load_index()
    target = next((i for i in items if i["id"] == upload_id), None)
    if not target:
        return False
    path = os.path.join(UPLOADS_DIR, target["disk_name"])
    if os.path.isfile(path):
        os.unlink(path)
    _save_index([i for i in items if i["id"] != upload_id])
    return True


def get_upload_path(upload_id: str) -> str | None:
    target = next((i for i in _load_index() if i["id"] == upload_id), None)
    if not target:
        return None
    path = os.path.realpath(os.path.join(UPLOADS_DIR, target["disk_name"]))
    if not path.startswith(os.path.realpath(UPLOADS_DIR)):
        return None
    return path if os.path.isfile(path) else None


def extract_text(upload_id: str, limit: int = 6000) -> str:
    """按文件类型提取文本内容，用于注入 LLM prompt"""
    path = get_upload_path(upload_id)
    if not path:
        return ""
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    try:
        if ext in TEXT_EXTS:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read(limit)
        if ext == ".docx":
            from docx import Document

            doc = Document(path)
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)[:limit]
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"【工作表：{ws.title}】")
                for row in ws.iter_rows(max_row=200, max_col=12, values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)[:limit]
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            text = "\n".join(
                (page.extract_text() or "") for page in reader.pages[:50]
            ).strip()
            if text:
                return text[:limit]
            return (
                f"（附件《{name}》未提取到可读文字，可能是扫描件/图片型 PDF。"
                "请把关键内容复制成文本，或另存为可复制文字的 PDF 后重试。）"
            )
        if ext in IMAGE_EXTS:
            return (
                f"（附件《{name}》是图片，已作为视觉输入随指令提交给模型识别；"
                "若成员模型不支持视觉，会自动忽略图片继续流程。）"
            )
        return f"（二进制文件：{name}，无法提取文本，请转换为 txt/docx/pdf 后上传）"
    except Exception as error:  # noqa: BLE001 —— 提取失败不影响主流程
        print(f"[Uploads] 提取失败 {path}：{error}")
        return f"（附件《{name}》解析失败：{str(error)[:80]}，请检查文件是否损坏）"


def extract_many(upload_ids: list[str], total_limit: int = 20000) -> str:
    """拼接多个附件的提取文本，带文件名标注（总长 2 万字，避免只看到开头一小段）"""
    if not upload_ids:
        return ""
    index = _load_index()
    parts: list[str] = []
    remaining = total_limit
    for uid in upload_ids:
        meta = next((i for i in index if i["id"] == uid), None)
        if not meta or remaining <= 0:
            continue
        text = extract_text(uid, limit=min(6000, remaining))
        if text:
            chunk = f"📎 附件《{meta['name']}》：\n{text}"
            parts.append(chunk)
            remaining -= len(chunk)
    return "\n\n".join(parts)


_IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
    ".tiff": "image/tiff",
}


def extract_images(
    upload_ids: list[str],
    max_count: int = 4,
    max_bytes: int = 2 * 1024 * 1024,
) -> list[str]:
    """把图片类附件读成 data URL（base64），供视觉模型多模态输入。

    最多 4 张、每张原始字节 ≤2MB；超限或非图片格式自动跳过，不阻断流程。"""
    if not upload_ids:
        return []
    import base64

    urls: list[str] = []
    for uid in upload_ids:
        if len(urls) >= max_count:
            break
        path = get_upload_path(uid)
        if not path:
            continue
        ext = os.path.splitext(path)[1].lower()
        if ext not in IMAGE_EXTS:
            continue
        try:
            if os.path.getsize(path) > max_bytes:
                print(f"[Uploads] 图片超过 2MB，跳过视觉输入：{os.path.basename(path)}")
                continue
            with open(path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode("ascii")
            urls.append(f"data:{_IMAGE_MIME.get(ext, 'image/png')};base64,{b64}")
        except OSError as error:
            print(f"[Uploads] 读取图片失败：{error}")
    return urls
