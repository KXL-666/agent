"""X1 source rows, editable formulas, cached results and chart range."""

import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from spreadsheet_builder import build_workbook
from spreadsheet_contract import audit_budget_records, budget_workbook
import office_runtime


X1 = ("界面设计：预算 10000 元，实际 9000 元。\n"
      "后端接口：预算 15000 元，实际 16500 元。\n"
      "测试验收：预算 8000 元，实际 7200 元。\n"
      "差额口径：预算减实际；执行率口径：实际除预算。")


class SpreadsheetContractTest(unittest.TestCase):
    def test_chart_cannot_include_total_even_with_explicit_end_row(self):
        sheets = budget_workbook(X1)
        sheets[0]["charts"][0]["end_row"] = 5
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            with self.assertRaisesRegex(ValueError, "包含汇总行"):
                build_workbook(sheets, str(Path(directory) / "invalid-total.xlsx"))

    def test_chart_rejects_text_data_column(self):
        sheets = budget_workbook(X1)
        sheets[0]["charts"][0]["data_columns"] = [1]
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            with self.assertRaisesRegex(ValueError, "非数值"):
                build_workbook(sheets, str(Path(directory) / "invalid-text.xlsx"))

    def test_source_audit_checks_amounts_not_only_labels(self):
        sheets = budget_workbook(X1)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = Path(directory) / "x1.xlsx"
            build_workbook(sheets, str(path))
            with closing(load_workbook(path, data_only=False)) as book:
                self.assertTrue(audit_budget_records(book, X1)["ok"])
                book.active["C3"] = 16501
                audit = audit_budget_records(book, X1)
                self.assertFalse(audit["ok"])
                self.assertEqual(audit["expected"], 3)
                self.assertEqual(audit["observed"], 3)

    def test_repeated_identical_source_records_are_not_silently_deduplicated(self):
        duplicate_source = ("研发接口：预算 1000 元，实际 900 元。\n"
                           "研发接口：预算 1000 元，实际 900 元。\n"
                           "预算实际差额执行率")
        sheets = budget_workbook(duplicate_source)
        self.assertIsNotNone(sheets)
        self.assertEqual(len(sheets[0]["rows"]), 3)  # 两笔相同记录和一行汇总
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = Path(directory) / "duplicates.xlsx"
            build_workbook(sheets, str(path))
            with closing(load_workbook(path, data_only=False)) as book:
                audit = audit_budget_records(book, duplicate_source)
                self.assertTrue(audit["ok"])
                self.assertEqual(audit["expected"], 2)
                self.assertEqual(audit["observed"], 2)

    def test_x1_real_workbook(self):
        sheets = budget_workbook(X1)
        self.assertEqual(len(sheets[0]["rows"]), 4)
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = Path(directory) / "x1.xlsx"
            self.assertTrue(build_workbook(sheets, str(path)))
            with closing(load_workbook(path, data_only=False)) as book:
                sheet = book.active
                self.assertEqual([sheet[f"A{i}"].value for i in (2, 3, 4)], ["界面设计", "后端接口", "测试验收"])
                self.assertEqual(sheet["D2"].value, "=B2-C2")
                self.assertEqual(sheet["E4"].value, "=IFERROR(C4/B4,0)")
                self.assertEqual(sheet["B5"].value, "=SUM(B2:B4)")
                self.assertEqual(sheet["B2"].comment.text.startswith("界面设计"), True)
                self.assertEqual(len(sheet._charts), 1)
                formula = sheet._charts[0].series[0].val.numRef.f
                self.assertIn("$B$2:$B$4", formula)
                self.assertNotIn("$B$5", formula)
            with closing(load_workbook(path, data_only=True)) as values:
                cached = values.active["D2"].value
                if office_runtime.find_office():
                    self.assertEqual(cached, 1000)
                else:
                    self.assertIsNone(cached)


if __name__ == "__main__":
    unittest.main()
