"""Located repairs change only their target section, slide, cell or file."""

import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

import candidate_revision
import forge


class RevisionTest(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.old = forge.WORKSPACE_DIR, forge.DELIVERABLES_DIR
        forge.WORKSPACE_DIR = str(Path(self.scratch.name) / "workspace")
        forge.DELIVERABLES_DIR = str(Path(self.scratch.name) / "deliverables")
        Path(forge.WORKSPACE_DIR).mkdir()
        Path(forge.DELIVERABLES_DIR).mkdir()

    def tearDown(self):
        forge.WORKSPACE_DIR, forge.DELIVERABLES_DIR = self.old
        self.scratch.cleanup()

    def _source(self, name):
        path = Path(forge.output_path(f"_process/source/{name}"))
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    def test_word_keeps_unrelated_paragraph(self):
        source = self._source("notice.docx")
        doc = Document()
        doc.add_paragraph("联调时间 14:00")
        doc.add_paragraph("其他安排保持不变")
        doc.save(source)
        revised, reason, _ = candidate_revision.apply("S1", "docx", "_process/source/notice.docx",
                                                        {"old_text": "14:00", "replacement": "15:00", "location": "第1段"})
        self.assertEqual(reason, "")
        self.assertEqual([p.text for p in Document(forge.output_path(revised)).paragraphs],
                         ["联调时间 15:00", "其他安排保持不变"])
        self.assertEqual(Document(source).paragraphs[0].text, "联调时间 14:00")

    def test_word_can_delete_one_unrequested_placeholder_paragraph(self):
        source = self._source("notice.docx")
        doc = Document()
        doc.add_paragraph("11 月 4 日联调")
        doc.add_paragraph("会议链接：待补充")
        doc.save(source)
        revised, reason, detail = candidate_revision.apply(
            "S1", "docx", "_process/source/notice.docx",
            {"operation": "delete", "old_text": "会议链接：待补充", "location": "第2段"},
        )
        self.assertEqual(reason, "")
        self.assertEqual(detail["operation"], "delete")
        self.assertEqual([p.text for p in Document(forge.output_path(revised)).paragraphs], ["11 月 4 日联调"])

    def test_ppt_keeps_other_slide(self):
        source = self._source("deck.pptx")
        deck = Presentation()
        for title in ("第一周", "第二周"):
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = title
        deck.save(source)
        revised, reason, _ = candidate_revision.apply("S1", "pptx", "_process/source/deck.pptx",
                                                        {"location": "第2页", "old_text": "第二周", "replacement": "第三周"})
        self.assertEqual(reason, "")
        result = Presentation(forge.output_path(revised))
        self.assertEqual([s.shapes.title.text for s in result.slides], ["第一周", "第三周"])

    def test_excel_keeps_other_cell(self):
        source = self._source("budget.xlsx")
        book = Workbook()
        book.active.title = "预算"
        book.active["A1"] = "界面设计"
        book.active["B1"] = 10000
        book.save(source)
        revised, reason, _ = candidate_revision.apply("S1", "xlsx", "_process/source/budget.xlsx",
                                                        {"location": "预算!B1", "replacement": 9000})
        self.assertEqual(reason, "")
        with closing(load_workbook(forge.output_path(revised))) as result:
            self.assertEqual(result.active["A1"].value, "界面设计")
            self.assertEqual(result.active["B1"].value, 9000)

    def test_code_keeps_other_file_and_tests(self):
        source = self._source("project.zip")
        source.write_bytes(b"placeholder")
        project = Path(self.scratch.name) / "project"
        project.mkdir()
        (project / "app.py").write_text("def convert(x):\n    return x\n", encoding="utf-8")
        (project / "test_app.py").write_text("from app import convert\nassert convert(1) == 1\n", encoding="utf-8")
        (project / "README.md").write_text("Run python app.py", encoding="utf-8")
        revised, reason, detail = candidate_revision.apply("S1", "code", "_process/source/project.zip",
                                                             {"path": "app.py", "replacement": "def convert(x):\n    return x + 1\n"},
                                                             project_dir=str(project))
        self.assertEqual(reason, "")
        self.assertIsNotNone(revised)
        copied = Path(detail["project_dir"])
        self.assertEqual((copied / "test_app.py").read_bytes(), (project / "test_app.py").read_bytes())
        self.assertEqual((project / "app.py").read_text(encoding="utf-8"), "def convert(x):\n    return x\n")


if __name__ == "__main__":
    unittest.main()
