"""Focused remaining-scope tests. All model/runtime work is mocked."""
import hashlib
import json
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from fastapi import HTTPException
import httpx

import artifacts
import candidate_revision
import forge
import main
import revision_code
import revision_diff
import revision_jobs
import revision_outputs as outputs
import revision_service as service
import storage
import test_revision_sessions as fixtures


class RevisionExtensionsTest(unittest.IsolatedAsyncioTestCase):
    setUp = fixtures.RevisionSessionTest.setUp
    asyncTearDown = fixtures.RevisionSessionTest.asyncTearDown

    def generate_file(self, kind, spec, filename, roster):
        name = 'generated-' + str(len(list(Path(forge.DELIVERABLES_DIR).rglob('*')))) + '.' + kind
        if kind == 'docx':
            doc = Document(); doc.add_paragraph('联调时间14:00'); doc.add_paragraph('其他安排保持不变')
            doc.save(forge.output_path(name))
        elif kind == 'pptx':
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = '联调汇报'
            deck.save(forge.output_path(name))
        else:
            Path(forge.output_path(name)).write_text('生成文件', encoding='utf-8')
        return name

    def pdf_file(self, source, filename):
        from pypdf import PdfWriter
        path = forge.safe_deliverable_path(source)
        name = 'derived-' + hashlib.sha256(Path(path).read_bytes()).hexdigest()[:12] + '.pdf'
        writer = PdfWriter(); writer.add_blank_page(width=300, height=400)
        writer.add_metadata({'/SourceHash': hashlib.sha256(Path(path).read_bytes()).hexdigest()})
        with open(forge.output_path(name), 'wb') as stream:
            writer.write(stream)
        return name

    async def accepted(self, artifact, sid):
        artifacts.update_artifact(artifact['id'], quality_status='passed')

    async def pair(self):
        session = storage.create_revision_conversation({'target_kinds': ['docx', 'pdf']})
        with patch.object(forge, 'forge_file', side_effect=self.generate_file), \
             patch.object(forge, 'pdf_from_docx', side_effect=self.pdf_file), \
             patch.object(forge, 'extract_preview', return_value='文件预览'), \
             patch.object(service, '_validate', side_effect=self.accepted):
            result = await outputs.generate(session, '生成通知Word与PDF', self.contract, 'pair')
        return session, result['artifact']

    async def test_initial_generation_binds_to_same_conversation_and_retry_is_idempotent(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://test') as client:
            created = await client.post('/api/revision/session', json={'target_kinds': ['docx']})
            self.assertEqual(created.status_code, 200)
            session = created.json()
            request = {'request_id': 'generate-1', 'expected_artifact_id': '', 'feedback': '生成通知Word', 'action': 'generate'}
            with patch.object(revision_jobs, 'dispatch'):
                queued = await client.post(f'/api/revision/session/{session["session_id"]}/message', json=request)
            with patch.object(service, 'evolve_contract', return_value=self.contract), \
                 patch.object(forge, 'forge_file', side_effect=self.generate_file), \
                 patch.object(service, '_validate', side_effect=self.accepted):
                await revision_jobs.run(queued.json()['operation']['id'])
            snapshot = revision_jobs.session_snapshot(session['session_id'])
            self.assertTrue(snapshot['artifact_id'])
            self.assertEqual(snapshot['root_id'], snapshot['artifact_id'])
            self.assertEqual(snapshot['history'][0]['action'], 'generate')
            with patch.object(revision_jobs, 'dispatch') as dispatch:
                repeated = await client.post(f'/api/revision/session/{session["session_id"]}/message', json=request)
                self.assertEqual(repeated.status_code, 202)
                self.assertEqual(repeated.json()['operation']['id'], queued.json()['operation']['id'])
                dispatch.assert_not_called()

    async def test_word_edit_refreshes_derived_pdf_and_rollback_restores_whole_bundle(self):
        _, anchor = await self.pair()
        old = outputs.output_snapshot(anchor)
        with patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '16:00'}), \
             patch.object(forge, 'pdf_from_docx', side_effect=self.pdf_file), \
             patch.object(service, '_validate', side_effect=self.accepted):
            changed = (await outputs.edit(anchor, '改为16:00', self.contract, operation_id='edit'))['artifact']
            revised = outputs.output_snapshot(changed)
            self.assertTrue(all(a['artifact_id'] != b['artifact_id'] for a, b in zip(old, revised)))
            word = next(o['artifact'] for o in revised if o['kind'] == 'docx')
            pdf = next(o['artifact'] for o in revised if o['kind'] == 'pdf')
            self.assertEqual(pdf['regen']['derived_source_sha256'], word['regen']['sha256'])
            restored = (await outputs.rollback(changed, anchor, 'restore'))['artifact']
        restored_outputs = outputs.output_snapshot(restored)
        restored_word = next(o['artifact'] for o in restored_outputs if o['kind'] == 'docx')
        self.assertIn('14:00', Document(forge.safe_deliverable_path(restored_word['file_name'])).paragraphs[0].text)
        self.assertEqual(len(restored_outputs), 2)
        self.assertEqual(restored['quality_status'], 'passed')

    async def test_derived_conversion_failure_does_not_advance_session(self):
        _, anchor = await self.pair()
        session = storage.create_revision_session(anchor['id'])
        op, _ = storage.begin_revision_operation(session['id'], 'failed-pdf', anchor['id'], 'revise',
                                                 {'feedback': '修改时间', 'mode': 'targeted'})
        with patch.object(service, 'evolve_contract', return_value=self.contract), \
             patch.object(main, 'chat_json', return_value={'old_text': '14:00', 'replacement': '17:00'}), \
             patch.object(forge, 'pdf_from_docx', return_value=None):
            await revision_jobs.run(op['id'])
        self.assertEqual(storage.get_revision_session(session['id'])['artifact_id'], anchor['id'])
        self.assertEqual(storage.get_revision_operation(op['id'])['status'], 'failed')

    async def test_add_ppt_then_convert_it_without_changing_original_word(self):
        with patch.object(forge, 'forge_file', side_effect=self.generate_file), \
             patch.object(service, '_validate', side_effect=self.accepted):
            added = (await outputs.edit(self.original, '增加汇报PPT', self.contract,
                operation_id='add', action='add', target_kind='pptx'))['artifact']
            bundle = outputs.output_snapshot(added)
            self.assertEqual({o['kind'] for o in bundle}, {'docx', 'pptx'})
            self.assertIn(self.original['id'], [o['artifact_id'] for o in bundle])
            converted = (await outputs.edit(added, '把演示改为HTML', self.contract,
                operation_id='convert', action='convert', target_kind='html'))['artifact']
        self.assertEqual({o['kind'] for o in outputs.output_snapshot(converted)}, {'docx', 'html'})
        self.assertEqual(Document(forge.safe_deliverable_path(self.original['file_name'])).paragraphs[0].text, '联调时间14:00')

    def test_ppt_native_font_and_fill(self):
        name = 'deck.pptx'
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = '标题'
        deck.save(forge.output_path(name))
        revised, reason, _ = candidate_revision.apply(self.sid, 'pptx', name,
            {'operation': 'format', 'page': 1, 'shape': 1, 'color': 'FF0000', 'fill': 'EEEEEE', 'font_size': 24, 'bold': True})
        self.assertFalse(reason)
        shape = Presentation(forge.safe_deliverable_path(revised)).slides[0].shapes[0]
        self.assertEqual(str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb), 'FF0000')
        self.assertEqual(str(shape.fill.fore_color.rgb), 'EEEEEE')

    def test_excel_range_format_preserves_formula_and_is_comparable(self):
        name = 'table.xlsx'
        workbook = Workbook(); sheet = workbook.active; sheet.title = '预算'
        sheet['A1'] = 10; sheet['B1'] = '=A1*2'; workbook.save(forge.output_path(name))
        original = artifacts.add_artifact('final', '预算', 3, '摘要', file_name=name, sprint_id=self.sid)
        revised, reason, _ = candidate_revision.apply(self.sid, 'xlsx', original['file_name'],
            {'operation': 'format', 'location': '预算!A1:B1', 'fill': 'FFDD00', 'bold': True, 'number_format': '0.00', 'column_width': 25})
        self.assertFalse(reason)
        opened = load_workbook(forge.safe_deliverable_path(revised))
        try:
            self.assertEqual(opened['预算']['B1'].value, '=A1*2')
        finally:
            opened.close()
        modified = artifacts.add_artifact('final', '预算2', 3, '摘要', file_name=revised,
                                        parent_id=original['id'], version=2, sprint_id=self.sid)
        diff = revision_diff.compare(modified['id'], original['id'])
        self.assertTrue(diff['complete']); self.assertGreater(diff['format_changes'], 0)
        self.assertEqual(diff['content_changes'], 0)

    def test_full_diff_finds_text_beyond_card_summary(self):
        doc = Document(); doc.add_paragraph('正文' * 9000); doc.add_paragraph('末尾未修改')
        doc.save(forge.output_path('long.docx'))
        original = artifacts.add_artifact('final', '长文', 3, '相同摘要', file_name='long.docx', sprint_id=self.sid)
        doc.paragraphs[-1].text = '末尾已经修改'; doc.save(forge.output_path('long-new.docx'))
        modified = artifacts.add_artifact('final', '长文2', 3, '相同摘要', file_name='long-new.docx',
                                        parent_id=original['id'], version=2, sprint_id=self.sid)
        result = revision_diff.compare(modified['id'], original['id'])
        self.assertIn('末尾已经修改', result['text_diff'])
        self.assertGreater(result['content_changes'], 0)

    def test_multifile_refactor_keeps_old_tree_and_updates_interfaces_together(self):
        source = Path(self.temp.name) / 'old-project'; source.mkdir()
        (source / 'core.py').write_text('def add(a,b): return a+b\n', encoding='utf-8')
        (source / 'main.py').write_text('from core import add\nprint(add(1,2))\n', encoding='utf-8')
        (source / 'test_core.py').write_text('from core import add\ndef test_add():\n assert add(1,2)==3\n', encoding='utf-8')
        (source / 'README.md').write_text('python main.py', encoding='utf-8')
        plan = {'runtime': 'python', 'summary': '统一接口', 'files': [
            {'path': 'core.py', 'action': 'write'}, {'path': 'main.py', 'action': 'write'}]}
        raw = '===FILE: core.py\ndef add(a,b): return a+b\ndef multiply(a,b): return a*b\n===END\n===FILE: main.py\nfrom core import multiply\nprint(multiply(2,3))\n===END'
        with patch.object(revision_code.llm, 'chat_json', return_value=plan), \
             patch.object(revision_code.llm, 'chat_text', return_value=raw), \
             patch.object(revision_code.project_agent, 'evaluate', return_value={'checks': [], 'status': 'unchecked'}):
            result = revision_code.build('添加乘法功能并修改入口', str(source))
        root = Path(result['regen']['project_dir'])
        self.assertIn('multiply', (root / 'core.py').read_text())
        self.assertIn('multiply', (root / 'main.py').read_text())
        self.assertNotIn('multiply', (source / 'main.py').read_text())
        self.assertTrue((root / 'test_core.py').exists())

    def test_refactor_cannot_delete_tests_or_escape_project(self):
        for item in ({'path': 'test_core.py', 'action': 'delete'}, {'path': '../escape.py', 'action': 'write'}):
            with patch.object(revision_code.llm, 'chat_json', return_value={'runtime': 'python', 'files': [item]}):
                with self.assertRaises(HTTPException):
                    revision_code.build('重构项目')


if __name__ == '__main__':
    unittest.main()
