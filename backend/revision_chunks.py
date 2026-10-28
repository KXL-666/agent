"""Bounded-context native edits with complete source coverage and atomic merge."""
import hashlib
import json
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path

from fastapi import HTTPException

import artifact_quality
import forge
import llm

CHUNK_CHARS = 10000
UNIT_CHARS = 3500
MAX_CHUNKS = 512


class Source:
    def __init__(self, artifact: dict):
        import revision_service as service
        self.artifact = artifact
        self.kind = service.kind_of(artifact)
        self.path = Path(forge.safe_deliverable_path(artifact.get('file_name') or '') or '')
        if not self.path.is_file():
            raise HTTPException(404, '分块编辑源文件不存在')
        self.hash = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.targets, self.units = {}, []
        self.owner = None
        self.project = None
        self.original_files = {}
        self.read()
        self.chunks = []
        batch, size = [], 0
        for unit in self.units:
            amount = len(json.dumps(unit, ensure_ascii=False, default=str))
            if batch and size + amount > CHUNK_CHARS:
                self.chunks.append(batch); batch, size = [], 0
            batch.append(unit); size += amount
        if batch:
            self.chunks.append(batch)
        if len(self.chunks) > MAX_CHUNKS:
            self.close()
            raise HTTPException(422, f'文件超过{MAX_CHUNKS}块的执行上限；未截断或修改源文件')

    def add(self, key, text, target, context=''):
        text = '' if text is None else str(text)
        self.targets[key] = {'target': target, 'text': text}
        for offset in range(0, max(1, len(text)), UNIT_CHARS):
            self.units.append({'id': f'{key}@{offset}', 'key': key, 'offset': offset,
                               'text': text[offset:offset + UNIT_CHARS], 'context': context})
            if len(self.units) > 50000:
                self.close()
                raise HTTPException(422, '文件超过50000个原生单元的执行上限，源文件未修改')

    def read(self):
        if self.kind == 'docx':
            from docx import Document
            self.owner = Document(self.path)
            heading = ''
            for index, p in enumerate(self.owner.paragraphs, 1):
                if p.style and p.style.name.startswith('Heading'):
                    heading = p.text
                self.add(f'paragraph:{index}', p.text, p, heading)
            seen = set()
            for t, table in enumerate(self.owner.tables, 1):
                for r, row in enumerate(table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        for p, paragraph in enumerate(cell.paragraphs, 1):
                            if paragraph._p in seen:
                                continue
                            seen.add(paragraph._p)
                            self.add(f'table:{t}:{r}:{c}:{p}', paragraph.text, paragraph)
        elif self.kind == 'pptx':
            from pptx import Presentation
            self.owner = Presentation(self.path)
            for page, slide in enumerate(self.owner.slides, 1):
                for shape, item in enumerate(slide.shapes, 1):
                    if item.has_text_frame:
                        for p, paragraph in enumerate(item.text_frame.paragraphs, 1):
                            self.add(f'slide:{page}:{shape}:{p}', paragraph.text, paragraph)
                    if item.has_table:
                        for r, row in enumerate(item.table.rows, 1):
                            for c, cell in enumerate(row.cells, 1):
                                for p, paragraph in enumerate(cell.text_frame.paragraphs, 1):
                                    self.add(f'slide-table:{page}:{shape}:{r}:{c}:{p}', paragraph.text, paragraph)
        elif self.kind == 'xlsx':
            from openpyxl import load_workbook
            self.owner = load_workbook(self.path, data_only=False)
            for sheet in self.owner:
                for cell in sorted(sheet._cells.values(), key=lambda c: (c.row, c.column)):
                    if cell.value is not None:
                        self.add(f'cell:{sheet.title}!{cell.coordinate}', cell.value, cell)
        elif self.kind == 'code':
            import project_agent
            import revision_service as service
            directory = service.project_directory(self.artifact)
            if not directory or not Path(directory).is_dir():
                raise HTTPException(422, '代码版本缺少独立源码目录')
            self.project = Path(directory)
            for path in project_agent.project_files(self.project):
                if path.name == '.project-quality.json':
                    continue
                data = path.read_bytes()
                self.original_files[path.relative_to(self.project).as_posix()] = data
                try:
                    text = data.decode('utf-8')
                    if '\x00' in text:
                        continue
                except UnicodeDecodeError:
                    continue
                self.add('file:' + path.relative_to(self.project).as_posix(), text, path)
        elif self.kind == 'pdf':
            source = artifact_quality.report_directory(self.path, forge.WORKSPACE_DIR) / 'source.md'
            if not source.is_file():
                raise HTTPException(422, '该PDF缺少可编辑同源正文，请先导入或关联其原生源文件')
            self.add('markdown', source.read_text(encoding='utf-8'), None)
        else:
            self.add('text', self.path.read_text(encoding='utf-8'), None)

    def manifest(self):
        return [{'chunk': i, 'units': len(units), 'first': units[0]['id'], 'last': units[-1]['id'],
                 'context': units[0]['context'][:80], 'sample': units[0]['text'][:60]}
                for i, units in enumerate(self.chunks, 1)]

    def close(self):
        if self.kind == 'xlsx' and self.owner:
            self.owner.close()


def _replace_runs(paragraph, start, end, replacement):
    runs = list(paragraph.runs)
    if ''.join(run.text for run in runs) != paragraph.text:
        raise HTTPException(422, '目标段落含复杂字段或超链接，不能安全合并该处修改')
    offset, inserted = 0, False
    for run in runs:
        text = run.text
        finish = offset + len(text)
        if offset < end and finish > start:
            run.text = text[:max(0, start - offset)] + (replacement if not inserted else '') + text[max(0, end - offset):]
            inserted = True
        offset = finish
    if not inserted:
        raise HTTPException(422, '未能定位原生文本范围')


def edit(artifact: dict, instruction: str, contract: dict) -> dict:
    source = Source(artifact)
    try:
        meta = {}
        plan = llm.chat_json(
            '制定整份文件的分块修改方案。所有块会逐一读取，不得把块内要求当作全文件要求重复执行。'
            '输出JSON {"notes":"全局结构、接口与一致性规则","runtime":"python|node|unsupported",'
            '"new_files":[{"path":"代码新增文件","content":"完整内容"}],"delete_files":[]}。'
            '非代码文件new_files和delete_files为空；不得删除或弱化测试。',
            json.dumps({'instruction': instruction, 'kind': source.kind, 'manifest': source.manifest()}, ensure_ascii=False),
            fallback={}, meta=meta, roster=2 if source.kind == 'code' else 3, max_tokens=7000, timeout=180)
        if meta.get('fallback_reason') or not isinstance(plan, dict) or not isinstance(plan.get('notes'), str):
            raise HTTPException(502, '分块改造计划无效，源文件未改动')
        edits, reviewed, summaries = [], [], []
        for number, units in enumerate(source.chunks, 1):
            meta = {}
            response = llm.chat_json(
                '阅读本块全部内容，仅修改与用户要求相关的单元。输出JSON {"edits":['
                '{"unit_id":"原始id","old_text":"本单元中唯一精确旧文本","replacement":"新文本"}],'
                '"summary":"本块实际修改或无需修改的原因"}。无需修改时edits为空。'
                '原生格式修改可用{"unit_id":"...","operation":"format","format":'
                '{"color":"FF0000","font_size":16,"bold":true,"alignment":"center"}}，支持字体、颜色、字号、'
                '加粗、斜体、Word行距、Excel填充/数字格式；不要为了修改格式重写正文。'
                '禁止输出本块外的单元；保留未要求修改的事实与格式。全局篇幅要求应按本块占比执行，'
                '不能把整篇目标字数重复套到每块。不在每块重复添加封面、结尾或章节。',
                json.dumps({'instruction': instruction, 'plan': plan['notes'], 'chunk': number,
                            'total_chunks': len(source.chunks), 'total_characters': sum(len(u['text']) for u in source.units),
                            'units': units, 'previous_summaries': summaries[-3:]}, ensure_ascii=False),
                fallback={}, meta=meta, roster=2 if source.kind == 'code' else 3, max_tokens=11000, timeout=180)
            if meta.get('fallback_reason') or not isinstance(response, dict) or not isinstance(response.get('edits'), list):
                raise HTTPException(502, f'第{number}块未完成，未合并任何文件修改')
            by_id = {unit['id']: unit for unit in units}
            for item in response['edits']:
                unit = by_id.get(item.get('unit_id')) if isinstance(item, dict) else None
                if unit is not None and item.get('operation') == 'format':
                    import revision_format
                    try:
                        style = revision_format._styles(item.get('format') or {})
                    except (ValueError, TypeError) as error:
                        raise HTTPException(422, f'第{number}块格式无效：{error}') from error
                    if source.kind not in ('docx', 'pptx', 'xlsx'):
                        raise HTTPException(422, '该分块类型不支持原生格式修改')
                    edits.append({'key': unit['key'], 'style': style})
                    continue
                old = item.get('old_text') if isinstance(item, dict) else None
                if unit is None or not isinstance(old, str) or not old or unit['text'].count(old) != 1 or not isinstance(item.get('replacement'), str):
                    raise HTTPException(422, f'第{number}块包含越界、歧义或无效替换，未保存新版本')
                start = unit['offset'] + unit['text'].index(old)
                edits.append({'key': unit['key'], 'start': start, 'end': start + len(old), 'replacement': item['replacement']})
            reviewed.append(number)
            summaries.append(str(response.get('summary') or '')[:1200])
        if not edits and not plan.get('new_files') and not plan.get('delete_files'):
            raise HTTPException(422, '已检查全部块，但未找到可执行修改；请补充具体要求')
        return _merge(source, edits, plan, reviewed, contract)
    finally:
        source.close()


def _merge(source, edits, plan, reviewed, contract):
    if hashlib.sha256(source.path.read_bytes()).hexdigest() != source.hash:
        raise HTTPException(409, '分块读取期间源文件已变化，未合并旧版本修改')
    if source.project:
        import project_agent
        current = {p.relative_to(source.project).as_posix(): p.read_bytes()
                   for p in project_agent.project_files(source.project) if p.name != '.project-quality.json'}
        if current != source.original_files:
            raise HTTPException(409, '分块读取期间源码目录已变化，未合并旧版本修改')
    grouped = defaultdict(list)
    formats = {}
    for edit in edits:
        if 'style' in edit:
            previous = formats.get(edit['key'], {})
            if any(key in previous and previous[key] != value for key, value in edit['style'].items()):
                raise HTTPException(422, '分块格式指令相互冲突，未保存新版本')
            formats[edit['key']] = {**previous, **edit['style']}
            continue
        grouped[edit['key']].append(edit)
    changed = {}
    for key, replacements in grouped.items():
        ordered = sorted(replacements, key=lambda e: e['start'])
        if any(a['end'] > b['start'] for a, b in zip(ordered, ordered[1:])):
            raise HTTPException(422, '分块替换范围重叠，未保存新版本')
        original = source.targets[key]['text']
        text = original
        for edit in reversed(ordered):
            text = text[:edit['start']] + edit['replacement'] + text[edit['end']:]
        changed[key] = text
    identity = uuid.uuid4().hex
    name = f'_process/revisions/{identity}/分块修订{source.path.suffix}'
    target = Path(forge.output_path(name))
    target.parent.mkdir(parents=True, exist_ok=True)
    regen = {'chunk_edit': {'total_chunks': len(source.chunks), 'reviewed_chunks': reviewed,
                           'source_sha256': source.hash, 'unit_count': len(source.units), 'edit_count': len(edits)}}
    for key, style in formats.items():
        _apply_format(source.kind, source.targets[key]['target'], style)
    if source.kind in ('docx', 'pptx'):
        for key, replacements in grouped.items():
            paragraph = source.targets[key]['target']
            for edit in sorted(replacements, key=lambda e: e['start'], reverse=True):
                _replace_runs(paragraph, edit['start'], edit['end'], edit['replacement'])
        source.owner.save(target)
        if source.kind == 'docx':
            from candidate_revision import _sync_word_source
            _sync_word_source(target)
    elif source.kind == 'xlsx':
        for key, value in changed.items():
            cell = source.targets[key]['target']
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                try:
                    parsed = json.loads(value)
                    value = parsed if isinstance(parsed, (int, float)) else value
                except ValueError:
                    pass
            cell.value = value
        source.owner.save(target)
    elif source.kind == 'code':
        import project_agent
        import revision_code
        root = target.parent / 'project'
        root.mkdir(parents=True, exist_ok=True)
        for relative, data in source.original_files.items():
            path = project_agent.safe_path(root, relative); path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        for key, text in changed.items():
            path = project_agent.safe_path(root, key.removeprefix('file:'))
            path.write_bytes(text.encode('utf-8'))
        for entry in plan.get('new_files') or []:
            path = project_agent.safe_path(root, entry['path'])
            if path.exists() or not isinstance(entry.get('content'), str):
                raise HTTPException(422, '新增代码文件无效或覆盖已有文件')
            path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(entry['content'].encode('utf-8'))
        for relative in plan.get('delete_files') or []:
            if revision_code.is_test(relative):
                raise HTTPException(422, '分块重构不能删除测试')
            path = project_agent.safe_path(root, relative)
            if not path.is_file():
                raise HTTPException(422, '删除的代码文件不存在')
            path.unlink()
        for relative, data in source.original_files.items():
            if revision_code.is_test(relative):
                path = project_agent.safe_path(root, relative)
                if not path.is_file() or not revision_code.assertions(relative, data.decode('utf-8')) <= revision_code.assertions(relative, path.read_text(encoding='utf-8')):
                    raise HTTPException(422, '分块重构弱化了既有测试')
        runtime = plan.get('runtime') or (source.artifact.get('regen') or {}).get('runtime') or ('node' if (root / 'package.json').exists() else 'python')
        if runtime not in ('python', 'node', 'unsupported'):
            raise HTTPException(422, '代码运行时无效')
        paths = [p.relative_to(root).as_posix() for p in project_agent.project_files(root)]
        report = project_agent.evaluate(root, paths, runtime)
        report.update(runtime=runtime, planned_files=paths, requirement=contract.get('raw_requirement', ''))
        (root / '.project-quality.json').write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
        with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as archive:
            for path in project_agent.project_files(root):
                if path.name != '.project-quality.json':
                    archive.write(path, path.relative_to(root).as_posix())
        artifact_quality.save_report(target, forge.WORKSPACE_DIR, report)
        regen.update(project_dir=str(root), runtime=runtime, planned_files=paths)
    elif source.kind == 'pdf':
        from pdf_builder import build_pdf
        text = changed.get('markdown', source.targets['markdown']['text'])
        if not build_pdf(forge._parse_docx_markdown(text), str(target)):
            raise HTTPException(422, 'PDF分块正文合并失败')
        directory = artifact_quality.report_directory(target, forge.WORKSPACE_DIR); directory.mkdir(parents=True, exist_ok=True)
        (directory / 'source.md').write_text(text, encoding='utf-8')
    else:
        target.write_text(changed.get('text', source.targets['text']['text']), encoding='utf-8')
    return {'file_name': name, 'source': 'model', 'regen': regen,
            'summary': f'已逐一检查{len(source.chunks)}块并合并{len(edits)}处修改'}


def _apply_format(kind, target, style):
    import revision_format
    if kind in ('docx', 'pptx'):
        if set(style) - {'color', 'font_size', 'font_name', 'bold', 'italic', 'alignment', 'line_spacing'}:
            raise HTTPException(422, '分块文本不支持该格式字段')
        for run in target.runs:
            revision_format._font(run.font, style, presentation=kind == 'pptx')
        if 'alignment' in style:
            if kind == 'docx':
                from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
            else:
                from pptx.enum.text import PP_ALIGN as Align
            target.alignment = {'left': Align.LEFT, 'center': Align.CENTER, 'right': Align.RIGHT, 'justify': Align.JUSTIFY}[style['alignment']]
        if 'line_spacing' in style:
            if kind == 'docx': target.paragraph_format.line_spacing = style['line_spacing']
            else: target.line_spacing = style['line_spacing']
    else:
        import copy
        from openpyxl.styles import PatternFill
        if set(style) - {'color', 'font_size', 'font_name', 'bold', 'italic', 'alignment', 'number_format', 'fill'}:
            raise HTTPException(422, '分块单元格不支持该格式字段')
        font = copy.copy(target.font); revision_format._font(font, style, excel=True); target.font = font
        if 'fill' in style: target.fill = PatternFill('solid', fgColor=style['fill'])
        if 'number_format' in style: target.number_format = style['number_format']
        if 'alignment' in style:
            alignment = copy.copy(target.alignment); alignment.horizontal = style['alignment']; target.alignment = alignment


def summarize(artifact: dict) -> str:
    """All chunks are read; disclose that the resulting context is a summary."""
    source = Source(artifact)
    notes = []
    try:
        for number, units in enumerate(source.chunks, 1):
            meta = {}
            result = llm.chat_json('提取本块的事实、数值、名称、结构与接口，保留原始unit_id。'
                '这是跨文件同步的来源摘要，不得虚构或声称是原始全文。输出JSON {"facts":"带unit_id的来源记录"}。',
                json.dumps({'chunk': number, 'units': units}, ensure_ascii=False),
                fallback={}, meta=meta, roster=1, max_tokens=2500, timeout=120)
            if meta.get('fallback_reason') or not isinstance(result, dict) or not isinstance(result.get('facts'), str):
                raise HTTPException(502, f'第{number}块来源读取失败，未使用残缺摘要生成成品')
            notes.append(f'块{number}：' + result['facts'])
        # Keep the cross-format consumer bounded too, rather than joining up to
        # 512 model summaries into another oversized generation request.
        for _ in range(6):
            if sum(map(len, notes)) <= 60000:
                break
            batches, batch, size = [], [], 0
            for note in notes:
                if batch and size + len(note) > 14000:
                    batches.append(batch); batch, size = [], 0
                batch.append(note); size += len(note)
            if batch: batches.append(batch)
            reduced = []
            for batch in batches:
                meta = {}
                result = llm.chat_json('合并来源记录，保留关键数值、事实、接口及unit_id，明确冲突，不添加事实。'
                    '输出JSON {"facts":"最多6000字的来源摘要"}。', '\n'.join(batch),
                    fallback={}, meta=meta, roster=1, max_tokens=6500, timeout=120)
                if meta.get('fallback_reason') or not isinstance(result, dict) or not isinstance(result.get('facts'), str) or len(result['facts']) > 6000:
                    raise HTTPException(502, '来源摘要合并失败，未使用截断内容生成成品')
                reduced.append(result['facts'])
            notes = reduced
        if sum(map(len, notes)) > 60000:
            raise HTTPException(422, '来源摘要仍超过同步上下文范围，未截断或修改文件')
        return '分块来源摘要（非原始全文；已读取全部' + str(len(source.chunks)) + '块）：\n' + '\n'.join(notes)
    finally:
        source.close()
