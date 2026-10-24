"""可编辑的工作簿：原生公式、图表、验证、格式与真实重算。"""
from datetime import datetime
import re

import office_runtime


def safe_formula(value: str) -> str:
    if not value.startswith("=") or len(value) > 2000:
        raise ValueError("无效公式")
    if re.search(r"\[|\]|\||(?:WEBSERVICE|HYPERLINK|RTD|CALL|EXEC|REGISTER\.ID)\s*\(", value, re.I):
        raise ValueError("公式包含外部链接或可执行调用")
    return value


def build_workbook(sheets: list[dict], output: str) -> bool:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    names = set()
    for index, sheet in enumerate(sheets):
        name = re.sub(r"[\\/?*\[\]:]", "", str(sheet["name"]))[:31] or f"工作表{index + 1}"
        if name.lower() in names:
            raise ValueError(f"工作表名称重复：{name}")
        names.add(name.lower())
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        headers, rows = sheet["headers"], sheet["rows"]
        ws.append(headers)
        formats, types = sheet.get("formats", {}), sheet.get("column_types", {})
        for row_index, row in enumerate(rows, 2):
            for column, raw in enumerate(row, 1):
                header = headers[column - 1]
                value = raw
                if isinstance(value, str):
                    if value.startswith("="):
                        value = safe_formula(value)
                    elif types.get(header) == "date" and value:
                        value = datetime.fromisoformat(value)
                    elif re.fullmatch(r"-?(?:\d+\.?\d*|\.\d+)", value):
                        # 编号和带前导零的数字不能被擅自改写。
                        if types.get(header) in ("number", "currency", "percent"):
                            value = float(value)
                cell = ws.cell(row_index, column, value)
                cell.number_format = formats.get(header) or {"date": "yyyy-mm-dd", "currency": '#,##0.00;[Red](#,##0.00)', "percent": "0.0%", "number": "#,##0.00"}.get(types.get(header), "General")
        for coordinate, formula in sheet.get("formulas", {}).items():
            if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,4}", coordinate):
                raise ValueError("公式单元格坐标无效")
            if ws[coordinate].column > len(headers) or ws[coordinate].row > len(rows) + 1:
                raise ValueError("公式坐标超出数据区域")
            ws[coordinate] = safe_formula(str(formula))
        for row in ws:
            for cell in row:
                cell.font = Font(name="微软雅黑", size=10, color="182230" if cell.data_type == "f" else "2455A4")
                cell.alignment = Alignment(vertical="center", horizontal="right" if cell.data_type in ("n", "f") else "left", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="E6EBF1"))
        for cell in ws[1]:
            cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for column, header in enumerate(headers, 1):
            samples = [str(header)] + [str(row[column - 1] or "") for row in rows[:60] if column <= len(row) and not str(row[column - 1]).startswith("=")]
            width = min(42, max(12, max(sum(2 if ord(c) > 127 else 1 for c in text) for text in samples) + 3))
            letter = get_column_letter(column)
            ws.column_dimensions[letter].width = width
            if types.get(header) in ("number", "currency", "percent"):
                ws.conditional_formatting.add(f"{letter}2:{letter}{len(rows) + 1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="B42318")))
        for row_index in range(2, len(rows) + 2):
            # 按最长换行估算行高，长文本不会只显示第一行。
            lines = max(1, max((len(str(c.value or "")) // max(8, int(ws.column_dimensions[c.column_letter].width / 1.7))) + 1 for c in ws[row_index]))
            ws.row_dimensions[row_index].height = min(120, max(24, lines * 15))
        for coordinate, note in sheet.get("notes", {}).items():
            if re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,4}", coordinate):
                ws[coordinate].comment = Comment(str(note)[:3000], "来源与口径")
        for validation in sheet.get("validations", [])[:12]:
            column = validation.get("column", "")
            if isinstance(column, int):
                column = get_column_letter(column)
            if not re.fullmatch(r"[A-Z]{1,3}", str(column)) or column_index_from_string(column) > len(headers):
                continue
            options = validation.get("options", [])
            if options and all(isinstance(v, str) and "," not in v and '"' not in v for v in options):
                values = ",".join(options)
                if len(values) > 250:
                    continue
                dv = DataValidation(type="list", formula1='"' + values + '"', allow_blank=True)
                dv.errorTitle, dv.error, dv.showErrorMessage = "请选择有效值", "请使用下拉列表中的选项", True
                ws.add_data_validation(dv)
                dv.add(f"{column}2:{column}{len(rows) + 1}")
        if len(set(map(str, headers))) == len(headers):
            table = Table(displayName=f"Data{index + 1}", ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for chart_index, config in enumerate(sheet.get("charts", [])[:3]):
            chart_type = config.get("type", "bar")
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type, BarChart)()
            columns = config.get("data_columns", [2])
            category = int(config.get("category_column", 1))
            if not 1 <= category <= len(headers):
                raise ValueError("图表分类列超出范围")
            detail_last_row = len(rows) + 1
            if rows and str(rows[-1][category - 1]).strip() in ("合计", "总计", "汇总", "Total", "TOTAL"):
                detail_last_row -= 1
            last_row = config.get("end_row", detail_last_row)
            if type(last_row) is not int or not 2 <= last_row <= detail_last_row:
                raise ValueError("图表数据行范围无效或包含汇总行")
            for column in columns:
                if type(column) is not int or not 1 <= column <= len(headers):
                    raise ValueError("图表数据列超出范围")
                if not all(ws.cell(row_index, column).data_type == "f" or
                           type(ws.cell(row_index, column).value) in (int, float)
                           for row_index in range(2, last_row + 1)):
                    raise ValueError("图表数据列包含非数值或空白单元格")
                chart.add_data(Reference(ws, min_col=column, max_col=column, min_row=1, max_row=last_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=category, max_col=category, min_row=2, max_row=last_row))
            chart.title = str(config.get("title", "数据对比"))[:100]
            chart.style = 10
            chart.width, chart.height = 18, 9
            # 图表在数据下方，避免遮挡可编辑单元格。
            ws.add_chart(chart, f"A{len(rows) + 4 + chart_index * 19}")
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape" if len(headers) > 5 else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:1"
        bottom = len(rows) + 1 + len(sheet.get("charts", [])[:3]) * 19 + (3 if sheet.get("charts") else 0)
        ws.print_area = f"A1:{get_column_letter(max(len(headers), 9 if sheet.get('charts') else 1))}{bottom}"
        ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    wb.save(output)
    wb.close()
    if any(isinstance(v, str) and v.startswith("=") for s in sheets for r in s["rows"] for v in r) or any(s.get("formulas") for s in sheets):
        office_runtime.recalculate(output)
    return True
