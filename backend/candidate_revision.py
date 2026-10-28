"""Apply one located repair to a copy of a candidate; preserve other parts."""

import hashlib
import json
from contextlib import closing
import re
import shutil
import uuid
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

import forge
import project_agent
import artifact_quality


def _destination(sprint_id: str, source: Path) -> Path:
    safe = re.sub(r"[^\w-]", "", sprint_id) or "current"
    relative = f"_process/revisions/{safe}/{uuid.uuid4().hex}/{source.name}"
    path = Path(forge.output_path(relative))
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _replace_once(texts: list[tuple[object, str]], old: str, new: str) -> bool:
    matches = [(target, value) for target, value in texts if old in value]
    if len(matches) != 1 or matches[0][1].count(old) != 1:
        return False
    target, value = matches[0]
    paragraphs = ([target] if hasattr(target, 'runs') else
                  list(target.text_frame.paragraphs) if hasattr(target, 'text_frame') else
                  list(target.paragraphs) if hasattr(target, 'paragraphs') else [])
    located = [p for p in paragraphs if old in p.text]
    if len(located) == 1:
        paragraph = located[0]
        runs = list(paragraph.runs)
        if ''.join(run.text for run in runs) == paragraph.text:
            start = paragraph.text.index(old)
            end = start + len(old)
            offset = 0
            inserted = False
            for run in runs:
                text = run.text
                run_end = offset + len(text)
                if offset < end and run_end > start:
                    prefix = text[:max(0, start - offset)]
                    suffix = text[max(0, end - offset):]
                    run.text = prefix + (new if not inserted else '') + suffix
                    inserted = True
                offset = run_end
            return inserted
    # Do not flatten rich text, fields or several paragraphs as a side effect
    # of a content edit. The explicit rewrite route can handle those cases.
    return False


def _sync_word_source(path: Path) -> None:
    """A subsequent whole-file rewrite must see the edited Word, not old Markdown."""
    document = Document(path)
    lines = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style else ''
        match = re.match(r'Heading ([1-6])', style)
        lines.append(('#' * int(match.group(1)) + ' ' if match else '') + paragraph.text)
    for table in document.tables:
        for row in table.rows:
            lines.append('| ' + ' | '.join(cell.text for cell in row.cells) + ' |')
    directory = artifact_quality.report_directory(path, forge.WORKSPACE_DIR)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / 'source.md').write_text('\n\n'.join(lines), encoding='utf-8')


def _rebuild_textual(kind: str, target: Path, spec: str, content: str) -> bool:
    """由同源 Markdown 重新渲染 Word/PDF。"""
    if kind == "pdf":
        from pdf_builder import build_pdf
        return bool(build_pdf(forge._parse_docx_markdown(content), str(target)))
    cover = bool(artifact_quality.contract(spec, "docx").get("cover"))
    return bool(forge._build_docx(content, str(target), cover=cover))


def _split_markdown(content: str) -> tuple[list[str], str]:
    """把同源 Markdown 拆成标题行与正文；字数口径只统计正文。"""
    lines = content.splitlines()
    heads = [line for line in lines if line.strip().startswith("#")]
    body = "\n".join(line for line in lines if not line.strip().startswith("#")).strip()
    return heads, body


def _llm_exact_fit(markdown: str, spec: str, target: int, roster: int | None) -> str:
    """LLM 按“还差多少字”的确定性反馈把正文精确重写到目标字数，最多 3 轮收敛。"""
    import llm
    heads, body = _split_markdown(markdown)
    current = artifact_quality.counted_length(forge._docx_prose(markdown))
    for _ in range(3):
        if current == target:
            break
        delta = target - current
        meta: dict = {}
        resp = llm.chat_json(
            "你是严谨的文字编辑。把正文精确改写到目标字数。"
            "字数统一口径：汉字 + 英文/数字词，不含标点、空格、换行与标题。"
            f"当前正文 {current} 字，目标恰好 {target} 字，需{'增加' if delta > 0 else '减少'} {abs(delta)} 字。"
            "只输出 JSON {\"body\":\"新正文\"}，body 只含正文段落（不含任何标题），段落间用换行分隔。"
            "内容须自然连贯、紧扣主题、信息完整；不得重复堆砌，不得出现字数自报或编辑说明，不得保留 Markdown 标记。",
            json.dumps({"target": target, "current": current, "headings": heads, "body": body},
                       ensure_ascii=False),
            fallback={}, roster=roster, meta=meta, max_tokens=6500, timeout=180,
        )
        new_body = str((resp or {}).get("body") or "").strip()
        if not new_body:
            break
        body = new_body
        markdown = "\n".join(heads + [body])
        current = artifact_quality.counted_length(forge._docx_prose(markdown))
    return markdown


def _refit_length_exact(kind: str, source: Path, target: Path, spec: str, roster: int | None) -> tuple[str | None, str, dict]:
    """字数精确返工：LLM 迭代重写同源正文，直到统一口径恰好等于目标字数。"""
    source_markdown = artifact_quality.report_directory(source, forge.WORKSPACE_DIR) / "source.md"
    if kind not in ("docx", "pdf") or not source_markdown.is_file():
        return None, "篇幅返工需要 Word/PDF 的同源正文", {}
    requested = artifact_quality.contract(spec, "docx")
    target_count = int(requested.get("count") or 0)
    if not target_count:
        return None, "缺少目标字数，无法精确返工", {}
    content = source_markdown.read_text(encoding="utf-8")
    before = artifact_quality.counted_length(forge._docx_prose(content))
    updated = _llm_exact_fit(content, spec, target_count, roster)
    after = artifact_quality.counted_length(forge._docx_prose(updated))
    if after == before:
        return None, f"篇幅 {before} 字未能改写收敛（模型未返回有效正文）", {"before": before, "after": after}
    if not _rebuild_textual(kind, target, spec, updated):
        return None, "篇幅返工重新渲染失败", {}
    target_quality = artifact_quality.report_directory(target, forge.WORKSPACE_DIR)
    target_quality.mkdir(parents=True, exist_ok=True)
    (target_quality / "source.md").write_text(updated, encoding="utf-8")
    detail = {"location": "全文", "before": before, "after": after, "target": target_count}
    exact = "已精确命中" if after == target_count else f"仍差 {abs(after - target_count)} 字"
    return "_process/" + target.relative_to(forge.WORKSPACE_DIR).as_posix(), \
        f"篇幅精确返工 {before}→{after} 字（目标 {target_count} 字，{exact}）", detail


def _refit_length(kind: str, source: Path, target: Path, spec: str) -> tuple[str | None, str, dict]:
    """整篇篇幅返工：按统一字数口径确定性压缩同源正文，不做逐段替换。"""
    source_markdown = artifact_quality.report_directory(source, forge.WORKSPACE_DIR) / "source.md"
    if kind not in ("docx", "pdf") or not source_markdown.is_file():
        return None, "篇幅返工需要 Word/PDF 的同源正文", {}
    content = source_markdown.read_text(encoding="utf-8")
    before = artifact_quality.counted_length(forge._docx_prose(content))
    updated = forge._fit_docx_length(forge._reconcile_pending_claims(content), spec)
    after = artifact_quality.counted_length(forge._docx_prose(updated))
    if after == before:
        return None, f"篇幅 {before} 字未能继续压缩（缺少可删减的完整句子）", {"before": before, "after": after}
    if kind == "pdf":
        from pdf_builder import build_pdf
        if not build_pdf(forge._parse_docx_markdown(updated), str(target)):
            return None, "PDF 篇幅返工生成失败", {}
    else:
        cover = artifact_quality.contract(spec, "docx")["cover"]
        if not forge._build_docx(updated, str(target), cover=cover):
            return None, "Word 篇幅返工生成失败", {}
    target_quality = artifact_quality.report_directory(target, forge.WORKSPACE_DIR)
    target_quality.mkdir(parents=True, exist_ok=True)
    (target_quality / "source.md").write_text(updated, encoding="utf-8")
    budget = artifact_quality.length_budget(spec, "docx") or {}
    detail = {"location": "全文", "before": before, "after": after, "target": budget.get("target")}
    return "_process/" + target.relative_to(forge.WORKSPACE_DIR).as_posix(), \
        f"篇幅按统一口径压缩 {before}→{after} 字（目标 {budget.get('target', '未指定')}）", detail


def apply(sprint_id: str, kind: str, source_file: str, patch: dict,
          project_dir: str = "", roster: int | None = None) -> tuple[str | None, str, dict]:
    source_name = forge.safe_deliverable_path(source_file)
    if not source_name:
        return None, "候选文件不存在", {}
    source = Path(source_name)
    target = _destination(sprint_id, source)
    location = str(patch.get("location") or "").strip()
    replacement = patch.get("replacement")
    operation = str(patch.get("operation") or "replace").strip().lower()
    if operation == 'format':
        import revision_format
        try:
            detail = revision_format.apply(kind, source, target, patch)
            return '_process/' + target.relative_to(forge.WORKSPACE_DIR).as_posix(), '', detail
        except (ValueError, IndexError, KeyError, TypeError) as error:
            return None, f'格式修订失败：{error}', {}
    if operation == "refit_length_exact":
        try:
            return _refit_length_exact(kind, source, target, str(patch.get("requirement") or ""), roster)
        except (OSError, RuntimeError, ValueError) as error:
            return None, f"篇幅精确返工失败：{error}", {}
    if operation == "refit_length":
        try:
            return _refit_length(kind, source, target, str(patch.get("requirement") or ""))
        except (OSError, RuntimeError, ValueError) as error:
            return None, f"篇幅返工失败：{error}", {}
    if operation == "delete" and kind != "docx":
        return None, "仅 Word 完整段落支持删除修订", {}
    if operation not in ("delete", "format") and (replacement is None or replacement == ""):
        return None, "缺少具体替换内容", {}
    detail = {}
    try:
        if kind == "docx":
            old = str(patch.get("old_text") or "")
            if operation != 'replace' and operation != 'delete':
                return None, 'Word 不支持此修改操作', {}
            if not old:
                return None, "Word 修订需要精确旧文本", {}
            doc = Document(source)
            if operation == "delete":
                matches = [p for p in doc.paragraphs if p.text.strip() == old.strip()]
                if len(matches) != 1 or len(doc.paragraphs) <= 1:
                    return None, "Word 删除需要唯一的完整段落，且不能删除全部正文", {}
                element = matches[0]._element
                element.getparent().remove(element)
            else:
                candidates = [(p, p.text) for p in doc.paragraphs]
                candidates += [(cell, cell.text) for table in doc.tables for row in table.rows for cell in row.cells]
                if not _replace_once(candidates, old, str(replacement)):
                    return None, "Word 目标文字不存在或匹配不唯一", {}
            doc.save(target)
            _sync_word_source(target)
            detail = {"location": location or old, "old_text": old, "operation": operation}
        elif kind == "pdf":
            old = str(patch.get("old_text") or "")
            source_markdown = artifact_quality.report_directory(source, forge.WORKSPACE_DIR) / "source.md"
            if not old or not source_markdown.is_file():
                return None, "PDF 定点修订需要同源 Markdown 与精确旧文本", {}
            content = source_markdown.read_text(encoding="utf-8")
            if content.count(old) != 1:
                return None, "PDF 源文本目标不存在或匹配不唯一", {}
            updated = content.replace(old, str(replacement), 1)
            from pdf_builder import build_pdf
            if not build_pdf(forge._parse_docx_markdown(updated), str(target)):
                return None, "PDF 修订版生成失败", {}
            target_quality = artifact_quality.report_directory(target, forge.WORKSPACE_DIR)
            target_quality.mkdir(parents=True, exist_ok=True)
            (target_quality / "source.md").write_text(updated, encoding="utf-8")
            detail = {"location": location or old, "old_text": old}
        elif kind == "pptx":
            match = re.search(r"(?:第\s*)?(\d+)\s*(?:页|张|slide)", location, re.I)
            old = str(patch.get("old_text") or "")
            if not match or not old:
                return None, "PPT 修订需要页码和精确旧文本", {}
            deck = Presentation(source)
            index = int(match.group(1)) - 1
            if index < 0 or index >= len(deck.slides):
                return None, "PPT 页码越界", {}
            candidates = [(shape, shape.text) for shape in deck.slides[index].shapes if shape.has_text_frame]
            if not _replace_once(candidates, old, str(replacement)):
                return None, "PPT 目标文字不存在或匹配不唯一", {}
            deck.save(target)
            detail = {"slide": index + 1, "old_text": old}
        elif kind == "xlsx":
            if "!" not in location:
                return None, "Excel 修订需要工作表!单元格位置", {}
            sheet_name, address = location.rsplit("!", 1)
            with closing(load_workbook(source)) as workbook:
                if sheet_name not in workbook or not re.fullmatch(r"\$?[A-Z]{1,3}\$?[1-9]\d*", address):
                    return None, "Excel 工作表或单元格位置无效", {}
                cell = workbook[sheet_name][address]
                old_value = cell.value
                cell.value = replacement
                workbook.save(target)
            detail = {"location": location, "old_value": old_value}
        elif kind == "code":
            relative = str(patch.get("path") or location).replace("\\", "/")
            if not project_dir or not relative:
                return None, "代码修订需要项目目录和文件路径", {}
            original_project = Path(project_dir).resolve()
            if not original_project.is_dir():
                return None, "集成项目源目录不存在", {}
            old_file = project_agent.safe_path(original_project, relative)
            if not old_file.is_file():
                return None, "代码目标文件不存在", {}
            project_copy = target.parent / "project"
            shutil.copytree(original_project, project_copy)
            new_file = project_agent.safe_path(project_copy, relative)
            old_content = new_file.read_text(encoding="utf-8")
            new_content = str(replacement)
            if new_file.name.startswith("test_") and new_content.count("assert") < old_content.count("assert"):
                return None, "返工不得减少测试断言", {}
            new_file.write_text(new_content, encoding="utf-8")
            checks = project_agent.static_check(project_copy)
            if any(check.get("ok") is False for check in checks):
                return None, "修订后的项目静态检查未通过", {"checks": checks}
            with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
                for path in project_agent.project_files(project_copy):
                    archive.write(path, path.relative_to(project_copy).as_posix())
            detail = {"path": relative, "project_dir": str(project_copy),
                      "old_sha256": hashlib.sha256(old_content.encode("utf-8")).hexdigest()}
        else:
            return None, f"{kind} 尚无安全的定点修订器", {}
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        return None, f"定点修订失败：{error}", {}
    return "_process/" + target.relative_to(forge.WORKSPACE_DIR).as_posix(), "", detail
